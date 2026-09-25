"""
End-to-end tests for the forced-substitution qualification diagnostic API.
Same setup pattern as test_api.py -- run with:

    pip install -r backend/requirements.txt
    python3 -m unittest backend.tests.test_change_case_api -v

Genuinely runtime-tested in this project's own sandbox during Phase 2
implementation (real FastAPI TestClient, real SQLite database, real
Alembic migrations applied) -- not a syntax-only port.
"""
import os
import shutil
import unittest


class TestChangeCaseApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.path.dirname(__file__), "_test_env_change_case")
        os.makedirs(cls.test_dir, exist_ok=True)
        os.environ["SECRET_KEY"] = "test-secret-key-not-for-production"
        os.environ["PASSWORD_PEPPER"] = "test-password-pepper-not-for-production"
        os.environ["DATABASE_URL"] = f"sqlite:///{cls.test_dir}/test.db"
        if os.path.exists(f"{cls.test_dir}/test.db"):
            os.remove(f"{cls.test_dir}/test.db")

        from fastapi.testclient import TestClient
        from backend.app.main import app
        from alembic.config import Config
        from alembic import command

        # init_db() (used by test_api.py) only runs the ORIGINAL
        # schema.sql and does not know about the Phase 2 tables, which
        # exist only as Alembic migrations -- see
        # PROJECT_ARCHITECTURE.md rule 6. Real Alembic migrations are the
        # only thing that creates the full current schema, so that's
        # what this test suite's setup must use.
        backend_dir = os.path.join(os.path.dirname(__file__), "..")
        alembic_cfg = Config(os.path.join(backend_dir, "alembic.ini"))
        alembic_cfg.set_main_option("script_location", os.path.join(backend_dir, "migrations"))
        command.upgrade(alembic_cfg, "head")

        cls.client_factory = staticmethod(lambda: TestClient(app))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.test_dir, ignore_errors=True)

    def setUp(self):
        self.client = self.client_factory()

    def _spec(self):
        return {
            "feature_columns": ["viscosity", "solids_pct"],
            "target_metric": "adhesion_score",
            "target_value": 80.0,
            "direction": "maximize",
        }

    def _signup(self, client, email, org_name):
        return client.post("/api/auth/signup", json={
            "email": email, "password": "correct-horse-battery-staple",
            "organization_name": org_name, "display_name": "Test Chemist",
        })

    def _create_change_case(self, client, name="PFHxA replacement"):
        return client.post("/api/change-cases", json={
            "name": name, "trigger_type": "regulatory_restriction",
            "restricted_substance": "PFHxA", "qualification_spec": self._spec(),
        })

    def _upgrade_org_to_paid(self, org_name: str):
        # Uses the app's own db_transaction() -- guaranteed to hit
        # whatever database the app itself is actually using -- rather
        # than a raw sqlite3.connect() to an assumed file path. The
        # latter broke when this test file ran in the same process as
        # test_api.py: `settings` is a singleton frozen at first import,
        # so a second test file's setUpClass re-assigning
        # os.environ["DATABASE_URL"] doesn't retroactively change which
        # database the already-imported app is connected to, while this
        # file's own self.db_path still pointed at its own (different,
        # unused-by-the-app) intended file. Real bug, found by running
        # the full combined suite, not just this file alone.
        from sqlalchemy import text
        from backend.app.config.database import db_transaction
        with db_transaction() as conn:
            conn.execute(
                text("UPDATE subscriptions SET plan='pilot' WHERE organization_id="
                     "(SELECT id FROM organizations WHERE name=:name)"),
                {"name": org_name},
            )

    def _csv(self, n=15):
        header = "viscosity,solids_pct,adhesion_score\n"
        rows = "\n".join(f"{450+i},{60+i*0.1},{70+i*0.5}" for i in range(n))
        return header + rows

    # -----------------------------------------------------------------
    # Core workflow
    # -----------------------------------------------------------------

    def test_full_diagnostic_workflow(self):
        self._signup(self.client, "workflow@coatings.com", "Workflow Coatings")
        r = self._create_change_case(self.client)
        self.assertEqual(r.status_code, 201)
        case_id = r.json()["id"]
        self.assertEqual(r.json()["status"], "draft")

        r = self.client.post(
            f"/api/change-cases/{case_id}/dataset",
            files={"file": ("history.csv", self._csv(), "text/csv")},
        )
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.json()["rows_ingested"], 15)

        r = self.client.post(f"/api/change-cases/{case_id}/candidates",
                              json={"name": "Resin A", "features": {"viscosity": 460, "solids_pct": 61}})
        self.assertEqual(r.status_code, 201)
        candidate_id = r.json()["id"]

        self._upgrade_org_to_paid("Workflow Coatings")

        r = self.client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200)
        results = r.json()
        self.assertEqual(len(results), 1)
        self.assertIn("predicted_probability", results[0])
        self.assertIn("recommended_experiment", results[0])

        r = self.client.post(
            f"/api/change-cases/{case_id}/candidates/{candidate_id}/outcome",
            json={"actual_result": "Passed adhesion at 84 vs target 80", "passed_spec": True},
        )
        self.assertEqual(r.status_code, 201)

    # -----------------------------------------------------------------
    # Entitlement gating
    # -----------------------------------------------------------------

    def test_ranking_blocked_on_trial_plan(self):
        client = self.client_factory()
        self._signup(client, "trial@coatings.com", "Trial Coatings")
        r = self._create_change_case(client)
        case_id = r.json()["id"]
        client.post(f"/api/change-cases/{case_id}/dataset",
                    files={"file": ("h.csv", self._csv(), "text/csv")})
        client.post(f"/api/change-cases/{case_id}/candidates",
                    json={"name": "A", "features": {"viscosity": 460, "solids_pct": 61}})

        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Trial", r.json()["error"])

    # -----------------------------------------------------------------
    # Tenant isolation -- every new endpoint, both trial and paid requester
    # -----------------------------------------------------------------

    def test_tenant_isolation_get_change_case(self):
        client_a = self.client_factory()
        self._signup(client_a, "iso-a1@coatings.com", "Iso Org A1")
        case_id = self._create_change_case(client_a).json()["id"]

        client_b = self.client_factory()
        self._signup(client_b, "iso-b1@coatings.com", "Iso Org B1")
        r = client_b.get(f"/api/change-cases/{case_id}")
        self.assertEqual(r.status_code, 404)

    def test_tenant_isolation_list_candidates(self):
        client_a = self.client_factory()
        self._signup(client_a, "iso-a2@coatings.com", "Iso Org A2")
        case_id = self._create_change_case(client_a).json()["id"]

        client_b = self.client_factory()
        self._signup(client_b, "iso-b2@coatings.com", "Iso Org B2")
        r = client_b.get(f"/api/change-cases/{case_id}/candidates")
        self.assertEqual(r.status_code, 404)

    def test_tenant_isolation_ranking_uniform_404_regardless_of_requester_plan(self):
        """The specific ordering bug found and fixed during Phase 2
        implementation: ranking must return 404 for a change case that
        isn't the requester's, whether the requester is on a trial plan
        or a paid one -- ownership must be checked before entitlement, so
        a trial-plan requester probing another org's change_case_id gets
        the same response as a paid one would."""
        client_a = self.client_factory()
        self._signup(client_a, "iso-a3@coatings.com", "Iso Org A3")
        case_id = self._create_change_case(client_a).json()["id"]

        # Trial requester
        client_b = self.client_factory()
        self._signup(client_b, "iso-b3@coatings.com", "Iso Org B3")
        r = client_b.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 404, "trial requester must get 404, not a 400 about their own plan")

        # Paid requester
        client_c = self.client_factory()
        self._signup(client_c, "iso-c3@coatings.com", "Iso Org C3")
        self._upgrade_org_to_paid("Iso Org C3")
        r = client_c.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 404)

    def test_tenant_isolation_record_outcome(self):
        client_a = self.client_factory()
        self._signup(client_a, "iso-a4@coatings.com", "Iso Org A4")
        case_id = self._create_change_case(client_a).json()["id"]
        client_a.post(f"/api/change-cases/{case_id}/dataset",
                       files={"file": ("h.csv", self._csv(), "text/csv")})
        cand = client_a.post(f"/api/change-cases/{case_id}/candidates",
                              json={"name": "A", "features": {"viscosity": 460, "solids_pct": 61}}).json()

        client_b = self.client_factory()
        self._signup(client_b, "iso-b4@coatings.com", "Iso Org B4")
        r = client_b.post(f"/api/change-cases/{case_id}/candidates/{cand['id']}/outcome",
                           json={"actual_result": "fabricated", "passed_spec": True})
        self.assertEqual(r.status_code, 404)

    # -----------------------------------------------------------------
    # Append-only outcomes, enforced at the route level
    # -----------------------------------------------------------------

    def test_duplicate_outcome_rejected_at_route_level(self):
        client = self.client_factory()
        self._signup(client, "append@coatings.com", "Append Coatings")
        case_id = self._create_change_case(client).json()["id"]
        client.post(f"/api/change-cases/{case_id}/dataset",
                    files={"file": ("h.csv", self._csv(), "text/csv")})
        cand = client.post(f"/api/change-cases/{case_id}/candidates",
                            json={"name": "A", "features": {"viscosity": 460, "solids_pct": 61}}).json()

        r1 = client.post(f"/api/change-cases/{case_id}/candidates/{cand['id']}/outcome",
                          json={"actual_result": "first", "passed_spec": True})
        self.assertEqual(r1.status_code, 201)

        r2 = client.post(f"/api/change-cases/{case_id}/candidates/{cand['id']}/outcome",
                          json={"actual_result": "second attempt", "passed_spec": False})
        self.assertEqual(r2.status_code, 400)

    # -----------------------------------------------------------------
    # Provenance: every prediction is tied to the dataset actually used
    # -----------------------------------------------------------------

    def test_prediction_provenance_tracks_dataset_version(self):
        client = self.client_factory()
        self._signup(client, "provenance@coatings.com", "Provenance Coatings")
        case_id = self._create_change_case(client).json()["id"]

        r = client.post(f"/api/change-cases/{case_id}/dataset",
                         files={"file": ("v1.csv", self._csv(15), "text/csv")})
        first_dataset_id = r.json()["dataset_id"]

        client.post(f"/api/change-cases/{case_id}/candidates",
                    json={"name": "A", "features": {"viscosity": 460, "solids_pct": 61}})
        self._upgrade_org_to_paid("Provenance Coatings")
        client.post(f"/api/change-cases/{case_id}/rank")

        candidates = client.get(f"/api/change-cases/{case_id}/candidates").json()
        self.assertEqual(candidates[0]["latest_prediction"]["dataset_version_id"], first_dataset_id)

        # Re-upload a NEW dataset version, re-rank, confirm the prediction
        # now points at the NEW dataset, not the stale first one.
        r2 = client.post(f"/api/change-cases/{case_id}/dataset",
                          files={"file": ("v2.csv", self._csv(20), "text/csv")})
        second_dataset_id = r2.json()["dataset_id"]
        self.assertNotEqual(first_dataset_id, second_dataset_id)

        client.post(f"/api/change-cases/{case_id}/rank")
        candidates = client.get(f"/api/change-cases/{case_id}/candidates").json()
        self.assertEqual(candidates[0]["latest_prediction"]["dataset_version_id"], second_dataset_id)

    # -----------------------------------------------------------------
    # Scope limits (3-5 candidates, no candidate discovery)
    # -----------------------------------------------------------------

    def test_candidate_limit_enforced(self):
        client = self.client_factory()
        self._signup(client, "limit@coatings.com", "Limit Coatings")
        case_id = self._create_change_case(client).json()["id"]

        for i in range(5):
            r = client.post(f"/api/change-cases/{case_id}/candidates",
                             json={"name": f"Candidate {i}", "features": {"viscosity": 460, "solids_pct": 61}})
            self.assertEqual(r.status_code, 201)

        r = client.post(f"/api/change-cases/{case_id}/candidates",
                         json={"name": "One too many", "features": {"viscosity": 460, "solids_pct": 61}})
        self.assertEqual(r.status_code, 400)

    # -----------------------------------------------------------------
    # Report generation -- the commercial deliverable
    # -----------------------------------------------------------------

    def test_report_generation_produces_real_docx(self):
        client = self.client_factory()
        self._signup(client, "report@coatings.com", "Report Coatings")
        case_id = self._create_change_case(client).json()["id"]
        client.post(f"/api/change-cases/{case_id}/dataset",
                    files={"file": ("h.csv", self._csv(), "text/csv")})
        client.post(f"/api/change-cases/{case_id}/candidates",
                    json={"name": "Resin A", "features": {"viscosity": 460, "solids_pct": 61}})
        self._upgrade_org_to_paid("Report Coatings")
        client.post(f"/api/change-cases/{case_id}/rank")

        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["content-type"],
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertGreater(len(r.content), 0)

    def test_tenant_isolation_report_generation(self):
        client_a = self.client_factory()
        self._signup(client_a, "iso-a5@coatings.com", "Iso Org A5")
        case_id = self._create_change_case(client_a).json()["id"]

        client_b = self.client_factory()
        self._signup(client_b, "iso-b5@coatings.com", "Iso Org B5")
        r = client_b.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 404)

    # -----------------------------------------------------------------
    # Priority 4 -- data-quality classification against REAL CSV uploads
    # (the rule function itself is unit-tested in test_change_case_rules.py;
    # these tests prove the ingestion service and the HTTP response
    # actually report it correctly.)
    # -----------------------------------------------------------------

    def _csv_from_rows(self, rows):
        """rows: list of (viscosity, solids_pct, adhesion_score) tuples."""
        header = "viscosity,solids_pct,adhesion_score\n"
        return header + "\n".join(f"{v},{s},{a}" for v, s, a in rows)

    def _upload_csv(self, client, case_id, csv_text):
        return client.post(
            f"/api/change-cases/{case_id}/dataset",
            files={"file": ("history.csv", csv_text, "text/csv")},
        )

    def _new_case(self, email, org_name):
        client = self.client_factory()
        self._signup(client, email, org_name)
        return client, self._create_change_case(client).json()["id"]

    def test_upload_clean_csv_reports_valid_quality(self):
        client, case_id = self._new_case("dq-clean@coatings.com", "DQ Clean Coatings")
        r = self._upload_csv(client, case_id, self._csv(15))
        self.assertEqual(r.status_code, 201)
        body = r.json()
        self.assertEqual(body["data_quality_status"], "valid")
        self.assertEqual(body["duplicate_rows_found"], 0)
        self.assertEqual(body["constant_columns"], [])

    def test_upload_detects_duplicate_rows_from_real_csv(self):
        # 15 distinct rows plus 2 exact repeats of the first row:
        # duplicates must be counted, but 15 distinct rows is still enough.
        rows = [(450 + i, 60 + i * 0.1, 70 + i * 0.5) for i in range(15)]
        rows += [rows[0], rows[0]]
        client, case_id = self._new_case("dq-dup@coatings.com", "DQ Dup Coatings")
        r = self._upload_csv(client, case_id, self._csv_from_rows(rows))
        self.assertEqual(r.status_code, 201)
        body = r.json()
        self.assertEqual(body["rows_ingested"], 17)
        self.assertEqual(body["duplicate_rows_found"], 2)
        self.assertEqual(body["data_quality_status"], "valid")

    def test_upload_many_duplicates_classified_insufficient(self):
        # 12 rows, but only 6 distinct -> fewer than the 8 distinct rows
        # required, so duplicates must not be allowed to inflate evidence.
        distinct = [(450 + i, 60 + i * 0.1, 70 + i * 0.5) for i in range(6)]
        client, case_id = self._new_case("dq-insuff@coatings.com", "DQ Insuff Coatings")
        r = self._upload_csv(client, case_id, self._csv_from_rows(distinct * 2))
        self.assertEqual(r.status_code, 201)
        body = r.json()
        self.assertEqual(body["rows_ingested"], 12)
        self.assertEqual(body["duplicate_rows_found"], 6)
        self.assertEqual(body["data_quality_status"], "insufficient")

    def test_upload_detects_constant_feature_column_from_real_csv(self):
        rows = [(450 + i, 60.0, 70 + i * 0.5) for i in range(15)]  # solids_pct never varies
        client, case_id = self._new_case("dq-const@coatings.com", "DQ Const Coatings")
        r = self._upload_csv(client, case_id, self._csv_from_rows(rows))
        self.assertEqual(r.status_code, 201)
        body = r.json()
        self.assertEqual(body["data_quality_status"], "invalid")
        self.assertEqual(body["constant_columns"], ["solids_pct"])

    def test_upload_small_distinct_csv_classified_insufficient(self):
        rows = [(450 + i, 60 + i * 0.1, 70 + i * 0.5) for i in range(5)]
        client, case_id = self._new_case("dq-small@coatings.com", "DQ Small Coatings")
        r = self._upload_csv(client, case_id, self._csv_from_rows(rows))
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.json()["data_quality_status"], "insufficient")

    # -----------------------------------------------------------------
    # Priority 4 -- bad data must not silently reach a ranking or report
    # -----------------------------------------------------------------

    def _case_with_dataset_and_candidate(self, email, org_name, csv_text):
        client, case_id = self._new_case(email, org_name)
        self.assertEqual(self._upload_csv(client, case_id, csv_text).status_code, 201)
        client.post(f"/api/change-cases/{case_id}/candidates",
                    json={"name": "Resin A", "features": {"viscosity": 460, "solids_pct": 61}})
        self._upgrade_org_to_paid(org_name)
        return client, case_id

    def test_ranking_blocked_when_dataset_has_constant_column(self):
        rows = [(450 + i, 60.0, 70 + i * 0.5) for i in range(15)]
        client, case_id = self._case_with_dataset_and_candidate(
            "gate-const@coatings.com", "Gate Const Coatings", self._csv_from_rows(rows))
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400)
        self.assertIn("data-quality", r.json()["error"])

    def test_ranking_blocked_when_distinct_rows_insufficient_due_to_duplicates(self):
        distinct = [(450 + i, 60 + i * 0.1, 70 + i * 0.5) for i in range(6)]
        client, case_id = self._case_with_dataset_and_candidate(
            "gate-dup@coatings.com", "Gate Dup Coatings", self._csv_from_rows(distinct * 2))
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400)

    def test_report_blocked_when_latest_dataset_invalid(self):
        rows = [(450 + i, 60.0, 70 + i * 0.5) for i in range(15)]
        client, case_id = self._case_with_dataset_and_candidate(
            "gate-report@coatings.com", "Gate Report Coatings", self._csv_from_rows(rows))
        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 400)
        self.assertIn("data-quality", r.json()["error"])

    # -----------------------------------------------------------------
    # Priority 4 -- report evidence basis and dataset provenance
    # -----------------------------------------------------------------

    INSUFFICIENT_STATEMENT = "Insufficient evidence to establish predictive performance."

    def _docx_text(self, content: bytes) -> str:
        import io
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    def _latest_prediction_dataset_id(self, client, case_id):
        candidates = client.get(f"/api/change-cases/{case_id}/candidates").json()
        return candidates[0]["latest_prediction"]["dataset_version_id"]

    def test_insufficient_dataset_can_generate_report_with_insufficient_evidence_language(self):
        distinct = [(450 + i, 60 + i * 0.1, 70 + i * 0.5) for i in range(6)]
        client, case_id = self._case_with_dataset_and_candidate(
            "rep-insuff@coatings.com", "Rep Insuff Coatings", self._csv_from_rows(distinct * 2))
        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 200)
        text = self._docx_text(r.content)
        self.assertIn(self.INSUFFICIENT_STATEMENT, text)
        self.assertIn("Data quality status: insufficient", text)
        # must not claim a model basis the data cannot support
        self.assertNotIn("Bayesian surrogate model", text)

    def test_clean_dataset_ranks_and_reports_normally(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "rep-clean@coatings.com", "Rep Clean Coatings", self._csv(15))
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200)
        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 200)
        text = self._docx_text(r.content)
        self.assertNotIn(self.INSUFFICIENT_STATEMENT, text)
        self.assertIn("Data quality status: valid", text)
        self.assertIn("Predictions are based on 15 historical qualification", text)
        self.assertIn(f"Dataset version: {self._latest_prediction_dataset_id(client, case_id)}", text)

    def test_older_prediction_and_report_stay_tied_to_original_dataset_after_new_upload(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "prov-old@coatings.com", "Prov Old Coatings", self._csv(15))
        d1 = self._latest_dataset_id(client, case_id)
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 200)
        self.assertEqual(self._latest_prediction_dataset_id(client, case_id), d1)

        # A NEWER upload that is invalid (20 rows, constant solids_pct).
        bad_rows = [(450 + i, 60.0, 70 + i * 0.5) for i in range(20)]
        r = self._upload_csv(client, case_id, self._csv_from_rows(bad_rows))
        d2 = r.json()["dataset_id"]
        self.assertNotEqual(d1, d2)
        self.assertEqual(r.json()["data_quality_status"], "invalid")

        # The existing prediction still points at D1, not the newer upload...
        self.assertEqual(self._latest_prediction_dataset_id(client, case_id), d1)

        # ...and the report is judged on D1 (valid, 15 rows), so it is NOT
        # blocked by the newer invalid D2, and it names D1, not D2.
        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 200)
        text = self._docx_text(r.content)
        self.assertIn(f"Dataset version: {d1}", text)
        self.assertNotIn(f"Dataset version: {d2}", text)
        self.assertIn("Predictions are based on 15 historical qualification", text)
        self.assertIn("Data quality status: valid", text)

        # A NEW ranking would use D2, so it is blocked -- and the failed
        # attempt does not silently move the existing prediction to D2.
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self._latest_prediction_dataset_id(client, case_id), d1)

    def test_reranking_after_valid_new_upload_creates_new_prediction_and_keeps_old(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "prov-new@coatings.com", "Prov New Coatings", self._csv(15))
        d1 = self._latest_dataset_id(client, case_id)
        client.post(f"/api/change-cases/{case_id}/rank")

        d2 = self._upload_csv(client, case_id, self._csv(18)).json()["dataset_id"]
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 200)
        self.assertEqual(self._latest_prediction_dataset_id(client, case_id), d2)

        from sqlalchemy import text
        from backend.app.config.database import db_connection
        with db_connection() as conn:
            stamped = [row[0] for row in conn.execute(
                text("SELECT dataset_version_id FROM predictions p JOIN candidate_substitutes c "
                     "ON c.id = p.candidate_id WHERE c.change_case_id = :cid ORDER BY p.id"),
                {"cid": case_id})]
        self.assertEqual(stamped, [d1, d2])  # the old prediction row is preserved, never rewritten

        text_out = self._docx_text(client.post(f"/api/change-cases/{case_id}/report").content)
        self.assertIn(f"Dataset version: {d2}", text_out)
        self.assertIn("Predictions are based on 18 historical qualification", text_out)

    def test_report_refuses_to_mix_predictions_from_different_dataset_versions(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "prov-mix@coatings.com", "Prov Mix Coatings", self._csv(15))
        client.post(f"/api/change-cases/{case_id}/candidates",
                    json={"name": "Resin B", "features": {"viscosity": 470, "solids_pct": 62}})
        client.post(f"/api/change-cases/{case_id}/rank")
        d2 = self._upload_csv(client, case_id, self._csv(18)).json()["dataset_id"]

        # Simulate a partial re-score: only one candidate gets a D2 prediction.
        from backend.app.repositories import candidate_repo
        org_id = self._org_id("Prov Mix Coatings")
        first = candidate_repo.list_candidates_for_change_case(org_id, case_id)[0]
        candidate_repo.create_prediction(org_id, first["id"], d2, "test-model", 0.5, 0.1)

        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 400)
        self.assertIn("different dataset versions", r.json()["error"])

    def _latest_dataset_id(self, client, case_id):
        from backend.app.repositories import qualification_dataset_repo
        org_id = self._org_id_for_case(case_id)
        return qualification_dataset_repo.get_latest_dataset_for_change_case(org_id, case_id)["id"]

    def _org_id(self, org_name):
        from sqlalchemy import text
        from backend.app.config.database import db_connection
        with db_connection() as conn:
            return conn.execute(text("SELECT id FROM organizations WHERE name = :n"), {"n": org_name}).scalar()

    def _org_id_for_case(self, case_id):
        from sqlalchemy import text
        from backend.app.config.database import db_connection
        with db_connection() as conn:
            return conn.execute(text("SELECT organization_id FROM change_cases WHERE id = :i"), {"i": case_id}).scalar()

    # -----------------------------------------------------------------
    # Priority 5 -- domain / extrapolation coverage (heuristic margin)
    # The default _csv(15) history spans viscosity 450..464 and
    # solids_pct 60.0..61.4, so with the 10% heuristic margin:
    #   viscosity near-edge up to 465.4, solids_pct up to 61.54.
    # -----------------------------------------------------------------

    def _add_candidate(self, client, case_id, name, viscosity, solids_pct):
        r = client.post(f"/api/change-cases/{case_id}/candidates",
                        json={"name": name, "features": {"viscosity": viscosity, "solids_pct": solids_pct}})
        self.assertEqual(r.status_code, 201)
        return r.json()["id"]

    def _rank_by_name(self, client, case_id):
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200)
        return {x["candidate_name"]: x for x in r.json()}

    def test_rank_reports_within_domain_for_in_range_candidate(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "dom-in@coatings.com", "Dom In Coatings", self._csv(15))
        res = self._rank_by_name(client, case_id)["Resin A"]
        cov = res["domain_coverage"]
        self.assertEqual(cov["status"], "within_historical_domain")
        self.assertEqual(cov["edge_margin_pct"], 0.10)
        self.assertEqual(cov["features"]["viscosity"]["historical_min"], 450)
        self.assertEqual(cov["features"]["viscosity"]["historical_max"], 464)
        self.assertIn("not a scientifically validated", cov["note"])
        self.assertNotIn("Domain caution", res["recommended_experiment"])

    def test_rank_flags_near_edge_and_outside_without_blocking_or_dropping(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "dom-flag@coatings.com", "Dom Flag Coatings", self._csv(15))
        self._add_candidate(client, case_id, "Edge", 465, 61)      # viscosity just past max, within 10%
        self._add_candidate(client, case_id, "Far", 600, 61)       # viscosity far outside
        self._add_candidate(client, case_id, "Mixed", 460, 70)     # one feature fine, one far outside
        res = self._rank_by_name(client, case_id)
        self.assertEqual(set(res), {"Resin A", "Edge", "Far", "Mixed"})  # warning, never a block or a drop

        self.assertEqual(res["Resin A"]["domain_coverage"]["status"], "within_historical_domain")
        self.assertEqual(res["Edge"]["domain_coverage"]["status"], "near_edge_of_domain")
        self.assertEqual(res["Far"]["domain_coverage"]["status"], "outside_historical_domain")
        mixed = res["Mixed"]["domain_coverage"]
        self.assertEqual(mixed["status"], "outside_historical_domain")  # worst feature wins
        self.assertEqual(mixed["features"]["viscosity"]["status"], "within_historical_domain")
        self.assertEqual(mixed["features"]["solids_pct"]["status"], "outside_historical_domain")

        # the stored recommended experiment carries the caution too, and names the flagged feature only
        self.assertIn("Domain caution", res["Far"]["recommended_experiment"])
        self.assertIn("outside the range", res["Far"]["recommended_experiment"])
        self.assertIn("near the edge of the range", res["Edge"]["recommended_experiment"])
        self.assertIn("data for solids_pct, so", res["Mixed"]["recommended_experiment"])
        self.assertNotIn("viscosity, solids_pct", res["Mixed"]["recommended_experiment"])

        # every candidate still got a stored, provenance-stamped prediction
        listed = client.get(f"/api/change-cases/{case_id}/candidates").json()
        self.assertEqual(len(listed), 4)
        for c in listed:
            self.assertIsNotNone(c["latest_prediction"])

    def test_candidate_listing_coverage_is_tied_to_prediction_dataset_not_latest_upload(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "dom-prov@coatings.com", "Dom Prov Coatings", self._csv(15))
        self._add_candidate(client, case_id, "Far", 600, 61)
        client.post(f"/api/change-cases/{case_id}/rank")

        def far_coverage():
            listed = {c["candidate_name"]: c for c in client.get(f"/api/change-cases/{case_id}/candidates").json()}
            return listed["Far"]["domain_coverage"]["status"]

        self.assertEqual(far_coverage(), "outside_historical_domain")

        # A newer, much WIDER dataset (viscosity 300..900) is uploaded. It must
        # not silently make the existing prediction look in-range.
        wide = [(300 + i * 40, 60 + i * 0.1, 70 + i * 0.5) for i in range(16)]
        self.assertEqual(self._upload_csv(client, case_id, self._csv_from_rows(wide)).status_code, 201)
        self.assertEqual(far_coverage(), "outside_historical_domain")

        # Only a NEW ranking on the wider dataset re-stamps the prediction --
        # and then coverage legitimately reflects that dataset.
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 200)
        self.assertEqual(far_coverage(), "within_historical_domain")

    def test_report_includes_coverage_column_warning_and_heuristic_disclaimer(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "dom-rep@coatings.com", "Dom Rep Coatings", self._csv(15))
        self._add_candidate(client, case_id, "Far", 600, 61)
        client.post(f"/api/change-cases/{case_id}/rank")
        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 200)
        text = self._docx_text(r.content)
        self.assertIn("Historical data coverage", text)
        self.assertIn("Within historical range", text)
        self.assertIn("Outside historical range", text)
        self.assertIn("prediction is an extrapolation beyond the historical data", text)
        self.assertIn("viscosity = 600", text)
        self.assertIn("heuristic", text)
        self.assertIn("not a scientifically validated", text)
        self.assertIn("10%", text)

    def test_report_for_in_range_candidates_has_no_extrapolation_warning(self):
        client, case_id = self._case_with_dataset_and_candidate(
            "dom-rep2@coatings.com", "Dom Rep2 Coatings", self._csv(15))
        client.post(f"/api/change-cases/{case_id}/rank")
        text = self._docx_text(client.post(f"/api/change-cases/{case_id}/report").content)
        self.assertIn("Within historical range", text)
        self.assertNotIn("extrapolation beyond the historical data", text)


if __name__ == "__main__":
    unittest.main()
