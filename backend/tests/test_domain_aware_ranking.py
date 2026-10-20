"""
Domain-aware ranking safeguard.

THE REPRODUCED DEFECT. Ranking was ordered purely by the model's estimated probability of meeting the
requirement. A candidate whose inputs lie far outside everything the historical data covers still gets a
probability from the same Gaussian-process model -- and because the model extrapolates its trend, that can be
an implausibly extreme prediction with probability 1.0. With the synthetic history used here (salt-spray
100-620 h observed) a candidate at ~3x the observed inputs was predicted at ~2,500 h, probability 1.0000, and
RANKED FIRST ahead of candidates the data actually supports. Ties at probability 1.0 were then broken in
favour of the LARGER predicted value, i.e. the wildest extrapolation won. The domain label existed
("Requires validation") but never influenced the order.

THE SAFEGUARD UNDER TEST. Candidates are ranked only against each other when the historical evidence supports
the comparison:
  * inputs within the observed range  -> "ranked"
  * mildly beyond it (heuristic margin) -> "ranked_with_caution" (ranked, flagged provisional)
  * clearly outside it / coverage unknown -> "not_ranked_insufficient_evidence": NO position (neither first
    nor last), the model's own numbers are reported UNCHANGED and labelled as extrapolation. Being outside the
    data is not evidence that a candidate is bad, so no probability is substituted and no penalty is applied.
  * if no candidate is supported, no candidate is ranked and the response says so.

WHAT THESE TESTS DO NOT SHOW. They use synthetic data. They prove the safeguard's logic and that the API, the
stored-prediction listing and the DOCX report agree; they do not show the model is scientifically valid or
that the heuristic range margin is a correct applicability domain.

Run:  python3 -m unittest backend.tests.test_domain_aware_ranking -v
"""
import io
import unittest

import numpy as np

from backend.tests import evidence_fixtures as F
from backend.tests._api_base import IsolatedApiTestCase

FEATURES = ["crosslinker_ratio", "cure_temp_c"]
# synthetic history: crosslinker 0.080-0.145, cure temp 140-179 C, salt spray 100-620 h  (14 distinct rows)
HISTORY = F.base_rows(14)
IN_A = {"crosslinker_ratio": 0.12, "cure_temp_c": 158}      # inside the observed range
IN_B = {"crosslinker_ratio": 0.10, "cure_temp_c": 150}      # inside the observed range
NEAR = {"crosslinker_ratio": 0.15, "cure_temp_c": 182}      # just past the maximum (within the 10% margin)
FAR = {"crosslinker_ratio": 0.40, "cure_temp_c": 300}       # ~3x the observed range
FAR_BELOW = {"crosslinker_ratio": 0.00, "cure_temp_c": 60}  # far outside, the other way

NOT_RANKED = "not_ranked_insufficient_evidence"


def _engine_estimates(history, candidates, target=300.0, direction="maximize"):
    """What the numerical engine says, called directly: the API must report exactly these numbers."""
    from engine.facade import rank_candidates
    X = np.array([[a, b] for a, b, _ in history])
    y = np.array([c for *_, c in history])
    out = rank_candidates(X, y, direction, FEATURES, target,
                          [{"id": i, "name": n, "features": f} for i, (n, f) in enumerate(candidates)])
    return {r["name"]: r for r in out}


class TestEvidenceRankingRules(unittest.TestCase):
    """The pure rule, with no database."""

    @staticmethod
    def _item(cid, prob, value, status):
        from backend.app.services import change_case_rules as R
        domain = {"within": R.DOMAIN_WITHIN, "near": R.DOMAIN_NEAR_EDGE, "out": R.DOMAIN_OUTSIDE, None: None}[status]
        coverage = {"status": domain} if domain else None
        return {"id": cid, "predicted_probability": prob, "predicted_value": value,
                "decision_support": R.derive_decision_support(coverage, has_prediction=True)}

    def _rank(self, items, direction="maximize"):
        from backend.app.services import ranking_evidence_rules as E
        return E.evidence_ranking(items, direction)

    def test_only_supported_candidates_get_positions_and_the_order_is_the_models(self):
        res = self._rank([
            self._item(1, 0.40, 380.0, "within"),
            self._item(2, 1.00, 2500.0, "out"),       # extreme extrapolation: must not take a position
            self._item(3, 0.90, 330.0, "near"),
            self._item(4, 0.95, 410.0, "within"),
        ])
        by = res["by_id"]
        self.assertIsNone(by[2]["rank"])
        self.assertEqual(by[2]["status"], NOT_RANKED)
        self.assertEqual([i for i in res["order"]][:3], [4, 3, 1])  # by model probability, as before
        self.assertEqual((by[4]["rank"], by[3]["rank"], by[1]["rank"]), (1, 2, 3))
        self.assertEqual((by[4]["status"], by[3]["status"], by[1]["status"]),
                         ("ranked", "ranked_with_caution", "ranked"))
        self.assertEqual(res["order"][-1], 2, "unranked candidates are listed after the ranked ones")

    def test_equal_probability_prefers_evidence_support_before_the_larger_extrapolated_value(self):
        """Probability saturates at 1.0, so many candidates tie. The old tie-break (bigger predicted value)
        rewarded extrapolation; evidence support now comes first."""
        res = self._rank([
            self._item(1, 1.0, 900.0, "near"),
            self._item(2, 1.0, 415.0, "within"),
        ])
        self.assertEqual(res["by_id"][2]["rank"], 1)
        self.assertEqual(res["by_id"][1]["rank"], 2)

    def test_probability_still_decides_between_unequal_candidates(self):
        res = self._rank([self._item(1, 0.99, 500.0, "near"), self._item(2, 0.60, 450.0, "within")])
        self.assertEqual((res["by_id"][1]["rank"], res["by_id"][2]["rank"]), (1, 2))

    def test_all_unsupported_means_nobody_is_ranked_and_the_summary_says_so(self):
        res = self._rank([self._item(1, 1.0, 2500.0, "out"), self._item(2, 0.0, -500.0, "out")])
        self.assertTrue(all(v["rank"] is None for v in res["by_id"].values()))
        self.assertEqual(res["summary"]["ranked"], 0)
        self.assertEqual(res["summary"]["not_ranked"], 2)
        self.assertIn("No candidate could be ranked", res["summary"]["message"])
        self.assertIn("does not mean", res["by_id"][2]["statement"])  # absence of a rank is not a verdict

    def test_unknown_coverage_fails_safe_to_not_ranked(self):
        res = self._rank([self._item(1, 0.9, 400.0, None), self._item(2, 0.5, 350.0, "within")])
        self.assertIsNone(res["by_id"][1]["rank"])
        self.assertEqual(res["by_id"][1]["status"], NOT_RANKED)
        self.assertEqual(res["by_id"][2]["rank"], 1)

    def test_a_single_supported_candidate_is_ranked_one(self):
        res = self._rank([self._item(7, 0.2, 310.0, "within")])
        self.assertEqual(res["by_id"][7]["rank"], 1)
        self.assertEqual(res["summary"]["message"], res["summary"]["message"])  # always present

    def test_missing_predicted_values_sort_after_recorded_ones_and_nothing_crashes(self):
        res = self._rank([self._item(1, 0.5, None, "within"), self._item(2, 0.5, 400.0, "within")])
        self.assertEqual((res["by_id"][2]["rank"], res["by_id"][1]["rank"]), (1, 2))

    def test_the_input_numbers_are_never_modified_and_input_order_does_not_matter(self):
        items = [self._item(1, 1.0, 2500.0, "out"), self._item(2, 0.5, 350.0, "within"),
                 self._item(3, 0.5, 350.0, "within"), self._item(4, 0.8, 360.0, "near")]
        snapshot = [dict(i) for i in items]
        first = self._rank(items)
        self.assertEqual(items, snapshot, "the rule must not rewrite probabilities, values or statuses")
        second = self._rank(list(reversed(items)))
        self.assertEqual({k: v["rank"] for k, v in first["by_id"].items()},
                         {k: v["rank"] for k, v in second["by_id"].items()})

    def test_direction_minimize_uses_the_lower_value_on_ties(self):
        res = self._rank([self._item(1, 1.0, 100.0, "within"), self._item(2, 1.0, 50.0, "within")], "minimize")
        self.assertEqual(res["by_id"][2]["rank"], 1)

    def test_every_status_has_a_statement_that_names_validation(self):
        res = self._rank([self._item(1, .9, 400., "within"), self._item(2, .8, 400., "near"), self._item(3, 1., 900., "out")])
        for v in res["by_id"].values():
            self.assertIn("validation", v["statement"].lower())
        self.assertIn("provisional", res["by_id"][2]["statement"].lower())
        self.assertIn("extrapolation", res["by_id"][3]["statement"].lower())


class TestDomainAwareRankingApi(IsolatedApiTestCase):
    dir_name = "_test_env_domain_rank"

    # ---------------------------------------------------------------- helpers
    def _case(self, label, history=HISTORY):
        org_id, client, case_id = self.setup_case(label)
        self.paid(org_id)
        up = self.upload(client, case_id, "history.csv", F.csv_bytes(F.CANON_HEADERS, history))
        self.assertEqual(up.status_code, 201, up.text)
        return org_id, client, case_id

    def _add(self, client, case_id, name, features):
        r = client.post(f"/api/change-cases/{case_id}/candidates", json={"name": name, "features": features})
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()["id"]

    def _rank(self, client, case_id):
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200, r.text)
        return {x["candidate_name"]: x for x in r.json()}, r.json()

    def _list(self, client, case_id):
        r = client.get(f"/api/change-cases/{case_id}/candidates")
        self.assertEqual(r.status_code, 200, r.text)
        return {c["candidate_name"]: c for c in r.json()}

    def _report(self, client, case_id):
        from docx import Document
        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 200, r.text)
        doc = Document(io.BytesIO(r.content))
        paragraphs = [p.text for p in doc.paragraphs]
        tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
        return paragraphs, tables

    @staticmethod
    def _table_with(tables, header_text):
        found = [t for t in tables if any(header_text in h for h in t[0])]
        return found[0] if found else None

    def _standard_case(self, label):
        org_id, client, case_id = self._case(label)
        ids = {n: self._add(client, case_id, n, f) for n, f in
               (("InA", IN_A), ("InB", IN_B), ("Near", NEAR), ("Far", FAR))}
        return org_id, client, case_id, ids

    # ---------------------------------------------------------------- the reproduced failure
    def test_reproduced_failure_a_far_out_candidate_must_not_rank_first(self):
        _, client, case_id, _ids = self._standard_case("repro")
        by, _ = self._rank(client, case_id)
        far = by["Far"]
        # the failure, stated as the facts that made it dangerous
        self.assertEqual(far["domain_coverage"]["status"], "outside_historical_domain")
        self.assertGreater(far["predicted_value"], 2 * max(c for *_, c in HISTORY),
                           "the synthetic extrapolation is implausibly high (the failure being reproduced)")
        self.assertEqual(far["predicted_probability"], 1.0)
        # the safeguard
        self.assertIsNone(far["rank"], "an out-of-domain candidate must not be given a position -- first or last")
        self.assertEqual(far["ranking_status"], NOT_RANKED)
        ranks = {n: r["rank"] for n, r in by.items() if r["rank"] is not None}
        self.assertEqual(sorted(ranks.values()), [1, 2, 3], "the supported candidates are ranked 1..n with no gap")
        self.assertNotIn("Far", ranks)

    def test_the_models_own_numbers_are_reported_unchanged_and_never_replaced(self):
        _, client, case_id, _ids = self._standard_case("unchanged")
        engine = _engine_estimates(HISTORY, [("InA", IN_A), ("InB", IN_B), ("Near", NEAR), ("Far", FAR)])
        by, _ = self._rank(client, case_id)
        for name, e in engine.items():
            self.assertAlmostEqual(by[name]["predicted_value"], e["predicted_value"], places=6, msg=name)
            self.assertAlmostEqual(by[name]["predicted_probability"], e["predicted_probability"], places=9, msg=name)
            self.assertAlmostEqual(by[name]["uncertainty_std"], e["uncertainty_std"], places=6, msg=name)
        far = by["Far"]
        self.assertNotIn(far["predicted_probability"], (0.0, 0.5), "no arbitrary probability may be substituted")
        # and what is stored is the same: the listing shows exactly what the run reported
        listed = self._list(client, case_id)["Far"]["latest_prediction"]
        self.assertAlmostEqual(listed["predicted_probability"], far["predicted_probability"], places=12)
        self.assertAlmostEqual(listed["predicted_value"], far["predicted_value"], places=9)

    def test_being_outside_the_data_is_not_treated_as_the_candidate_being_bad(self):
        """The unranked candidate must not be pushed to the bottom either: no rank at all, and the wording
        says the evidence cannot decide in either direction."""
        _, client, case_id, _ids = self._standard_case("notbad")
        by, _ = self._rank(client, case_id)
        far = by["Far"]
        self.assertIsNone(far["rank"])
        text = far["ranking_statement"].lower()
        self.assertIn("does not mean the candidate is unsuitable", text)
        self.assertIn("extrapolation", text)
        self.assertIn("physical validation", text)
        # in-domain candidates with a poor estimate ARE ranked (last), because the data supports that comparison
        self.assertEqual(by["InB"]["ranking_status"], "ranked")
        self.assertIsNotNone(by["InB"]["rank"])

    def test_near_edge_candidate_is_ranked_but_flagged_provisional(self):
        _, client, case_id, _ids = self._standard_case("near")
        by, _ = self._rank(client, case_id)
        near = by["Near"]
        self.assertEqual(near["domain_coverage"]["status"], "near_edge_of_domain")
        self.assertEqual(near["ranking_status"], "ranked_with_caution")
        self.assertIsNotNone(near["rank"])
        self.assertIn("provisional", near["ranking_statement"].lower())

    # ---------------------------------------------------------------- API consistency
    def test_run_and_stored_listing_agree_on_rank_status_and_statement(self):
        _, client, case_id, _ids = self._standard_case("agree")
        by, _ = self._rank(client, case_id)
        listed = self._list(client, case_id)
        for name, r in by.items():
            self.assertEqual(listed[name]["rank"], r["rank"], name)
            self.assertEqual(listed[name]["ranking_status"], r["ranking_status"], name)
            self.assertEqual(listed[name]["ranking_statement"], r["ranking_statement"], name)
            self.assertEqual(listed[name]["decision_support"]["status"], r["decision_support"]["status"], name)

    def test_run_response_is_ordered_ranked_first_then_unranked_and_carries_a_summary(self):
        _, client, case_id, _ids = self._standard_case("shape")
        _, rows = self._rank(client, case_id)
        statuses = [r["ranking_status"] for r in rows]
        self.assertEqual(statuses[-1], NOT_RANKED)
        ranked_part = [r["rank"] for r in rows if r["rank"] is not None]
        self.assertEqual(ranked_part, sorted(ranked_part))
        summaries = {str(r["ranking_summary"]) for r in rows}
        self.assertEqual(len(summaries), 1, "one summary, repeated on every result like model_quality")
        s = rows[0]["ranking_summary"]
        self.assertEqual((s["candidates"], s["ranked"], s["ranked_with_caution"], s["not_ranked"]), (4, 2, 1, 1))
        self.assertIn("1 of 4", s["message"])
        self.assertIn("ordering_rule", s)

    def test_rank_result_schema_accepts_an_unranked_candidate(self):
        from backend.app.schemas.change_case_schemas import RankResultResponse
        r = RankResultResponse(rank=None, candidate_id=1, candidate_name="x", predicted_probability=1.0,
                               uncertainty_std=1.0, recommended_experiment="e", ranking_status=NOT_RANKED,
                               ranking_statement="s")
        self.assertIsNone(r.rank)
        self.assertEqual(r.ranking_status, NOT_RANKED)

    # ---------------------------------------------------------------- everyone insufficiently supported
    def test_when_every_candidate_is_outside_the_evidence_nobody_is_ranked_and_it_is_said_plainly(self):
        org_id, client, case_id = self._case("allout")
        self._add(client, case_id, "FarA", FAR)
        self._add(client, case_id, "FarB", FAR_BELOW)
        by, rows = self._rank(client, case_id)
        self.assertTrue(all(r["rank"] is None and r["ranking_status"] == NOT_RANKED for r in rows))
        s = rows[0]["ranking_summary"]
        self.assertEqual((s["ranked"], s["ranked_with_caution"], s["not_ranked"]), (0, 0, 2))
        self.assertIn("No candidate could be ranked", s["message"])
        self.assertTrue(all(r["recommended_experiment"] for r in rows), "the validation plan is still produced")
        listed = self._list(client, case_id)
        self.assertTrue(all(c["rank"] is None for c in listed.values()))
        self.assertEqual({c["ranking_status"] for c in listed.values()}, {NOT_RANKED})
        # the run was still recorded as history
        self.assertTrue(all(c["prediction_count"] == 1 for c in listed.values()))

    def test_report_when_every_candidate_is_outside_the_evidence(self):
        _, client, case_id = self._case("allout_report")
        self._add(client, case_id, "FarA", FAR)
        self._add(client, case_id, "FarB", FAR_BELOW)
        self._rank(client, case_id)
        paragraphs, tables = self._report(client, case_id)
        text = " ".join(paragraphs)
        self.assertIn("No candidate could be ranked", text)
        self.assertIsNone(self._table_with(tables, "Rank"), "no ranked table may be shown when nobody is ranked")
        unranked = self._table_with(tables, "extrapolation")
        self.assertIsNotNone(unranked)
        self.assertEqual({row[0] for row in unranked[1:]}, {"FarA", "FarB"})
        self.assertIn("Recommended Validation Experiments", paragraphs)

    # ---------------------------------------------------------------- the DOCX report
    def test_report_ranks_only_supported_candidates_and_discloses_the_rest_separately(self):
        _, client, case_id, _ids = self._standard_case("report")
        by, _ = self._rank(client, case_id)
        paragraphs, tables = self._report(client, case_id)
        ranked = self._table_with(tables, "Rank")
        self.assertIsNotNone(ranked)
        self.assertEqual(ranked[0][0], "Rank")
        ranked_names = [row[1] for row in ranked[1:]]
        self.assertNotIn("Far", ranked_names, "an unsupported candidate must not appear in the ranked table")
        self.assertEqual(ranked_names, sorted(ranked_names, key=lambda n: by[n]["rank"]),
                         "the report orders exactly as the API does")
        self.assertEqual([row[0] for row in ranked[1:]], [str(by[n]["rank"]) for n in ranked_names])
        unranked = self._table_with(tables, "extrapolation")
        self.assertIsNotNone(unranked)
        self.assertEqual([row[0] for row in unranked[1:]], ["Far"])
        self.assertIn("Not ranked: outside the historical evidence", paragraphs)
        joined = " ".join(paragraphs)
        self.assertIn("not ranked", joined.lower())
        self.assertIn("does not mean", joined)
        # the disclosed number is the model's own, not a substitute
        self.assertTrue(any(str(round(by["Far"]["predicted_value"]))[:3] in cell for cell in unranked[1]),
                        unranked[1])
        # evidence gaps and the validation plan still cover the unranked candidate
        self.assertTrue(any(p.startswith("Far: ") and "not ranked" in p.lower() for p in paragraphs), paragraphs)
        self.assertTrue(any(p.startswith("Far: ") for p in paragraphs))

    def test_report_still_states_uncertainty_domain_and_validation_for_every_candidate(self):
        _, client, case_id, _ids = self._standard_case("report_cols")
        self._rank(client, case_id)
        _, tables = self._report(client, case_id)
        ranked = self._table_with(tables, "Rank")
        header = " | ".join(ranked[0])
        for needed in ("Uncertainty", "Historical data coverage", "Decision support", "Model-estimated probability"):
            self.assertIn(needed, header)
        unranked = self._table_with(tables, "extrapolation")
        header = " | ".join(unranked[0])
        for needed in ("Uncertainty", "Historical data coverage", "Decision support"):
            self.assertIn(needed, header)
        self.assertEqual(unranked[1][-1], "Requires validation")

    # ---------------------------------------------------------------- stale predictions
    def test_stale_predictions_are_never_presented_as_current_rankings(self):
        """Evidence is replaced AFTER the ranking ran: every earlier prediction is flagged stale on the
        listing, in the report, and a re-run makes it current again -- with the wider data the formerly
        out-of-domain candidate is now supported and ranked."""
        org_id, client, case_id, _ids = self._standard_case("stale")
        by, _ = self._rank(client, case_id)
        self.assertIsNone(by["Far"]["rank"])
        wider = [(round(0.08 + 0.005 * i, 4), 140 + 3 * i, 100 + 40 * i) for i in range(65)]  # reaches 0.40 / 332 C
        self.assertEqual(self.upload(client, case_id, "wider.csv", F.csv_bytes(F.CANON_HEADERS, wider)).status_code, 201)

        listed = self._list(client, case_id)
        for name, c in listed.items():
            self.assertIs(c["prediction_stale"], True, name)
            self.assertTrue(c["stale_notice"], name)
        # judged against the evidence it was computed from, the extrapolated candidate is STILL not ranked
        self.assertIsNone(listed["Far"]["rank"])
        self.assertEqual(listed["Far"]["ranking_status"], NOT_RANKED)
        paragraphs, _ = self._report(client, case_id)
        self.assertTrue(any("NEWER EVIDENCE EXISTS" in p.upper() or "newer evidence" in p.lower() for p in paragraphs),
                        "the report must say the predictions are from superseded evidence")

        by2, _ = self._rank(client, case_id)  # re-run on the wider evidence
        listed = self._list(client, case_id)
        for name, c in listed.items():
            self.assertIs(c["prediction_stale"], False, name)
            self.assertIsNone(c["stale_notice"], name)
        self.assertEqual(by2["Far"]["domain_coverage"]["status"], "within_historical_domain")
        self.assertIsNotNone(by2["Far"]["rank"], "now supported by evidence, so it is ranked on its merits")
        self.assertEqual(sorted(r["rank"] for r in by2.values() if r["rank"] is not None), [1, 2, 3, 4])

    def test_order_is_withheld_after_a_candidate_is_added_even_when_one_is_unranked(self):
        _, client, case_id, _ids = self._standard_case("outdated")
        self._rank(client, case_id)
        self._add(client, case_id, "Latecomer", {"crosslinker_ratio": 0.11, "cure_temp_c": 155})
        listed = self._list(client, case_id)
        self.assertTrue(all(c["rank"] is None for c in listed.values()), "ranks withheld until re-run")
        # ...but the evidence status stays visible, so nothing about the extrapolation is hidden
        self.assertEqual(listed["Far"]["ranking_status"], NOT_RANKED)
        self.assertEqual(listed["InA"]["ranking_status"], "ranked")
        self.assertTrue(listed["InA"]["ranking_notice"])

    # ---------------------------------------------------------------- not tied to the perfectly linear fixture
    def test_the_failure_and_the_safeguard_also_hold_on_noisy_history(self):
        noise = [12, -9, 15, -14, 7, -11, 13, -6, 10, -15, 9, -8, 14, -12]
        noisy = [(a, b, c + n) for (a, b, c), n in zip(HISTORY, noise)]
        _, client, case_id = self._case("noisy", history=noisy)
        for n, f in (("InA", IN_A), ("InB", IN_B), ("Far", FAR)):
            self._add(client, case_id, n, f)
        by, _ = self._rank(client, case_id)
        self.assertEqual(by["Far"]["domain_coverage"]["status"], "outside_historical_domain")
        self.assertIsNone(by["Far"]["rank"])
        self.assertEqual(sorted(r["rank"] for r in by.values() if r["rank"] is not None), [1, 2])
        engine = _engine_estimates(noisy, [("InA", IN_A), ("InB", IN_B), ("Far", FAR)])
        self.assertAlmostEqual(by["Far"]["predicted_probability"], engine["Far"]["predicted_probability"], places=9)


if __name__ == "__main__":
    unittest.main()
