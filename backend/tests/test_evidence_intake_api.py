"""
Step 4 -- adversarial evidence-intake tests through the REAL production path:
HTTP upload route -> evidence gate -> stored rows -> ranking -> report.
For every case the system must either handle the data correctly or explicitly
flag/reject it, and a REJECTED upload must leave NOTHING behind (no dataset row,
no experiment rows) so bad evidence can never reach the model.

Run:  python3 -m unittest backend.tests.test_evidence_intake_api -v
Own isolated SQLite database (rebind-and-assert pattern from test_api.py).
"""
import hashlib
import io
import json
import os
import shutil
import unittest
from contextlib import redirect_stdout

SPEC = {"feature_columns": ["crosslinker_ratio", "cure_temp_c"], "target_metric": "salt_spray_hours",
        "target_value": 300.0, "direction": "maximize"}
SPEC_H = {**SPEC, "units": {"salt_spray_hours": "h"}}
CANON = ["crosslinker_ratio", "cure_temp_c", "salt_spray_hours"]


def base_rows(n=12):
    return [(round(0.08 + 0.005 * i, 4), 140 + 3 * i, 100 + 40 * i) for i in range(n)]


def make_csv(header, rows):
    return ",".join(header) + "\n" + "\n".join(",".join(str(c) for c in r) for r in rows) + "\n"


class TestEvidenceIntakeApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.path.dirname(__file__), "_test_env_intake")
        os.makedirs(cls.test_dir, exist_ok=True)
        os.environ["SECRET_KEY"] = "test-secret-key-not-for-production"
        os.environ["PASSWORD_PEPPER"] = "test-password-pepper-not-for-production"
        os.environ["DATABASE_URL"] = f"sqlite:///{cls.test_dir}/test.db"
        if os.path.exists(f"{cls.test_dir}/test.db"):
            os.remove(f"{cls.test_dir}/test.db")
        from backend.app.config.settings import settings
        from backend.app.config import database
        settings.DATABASE_URL = os.environ["DATABASE_URL"]
        database.engine = database._make_engine()
        assert str(database.engine.url).endswith(f"{cls.test_dir}/test.db"), (
            f"Test isolation failure: engine bound to {database.engine.url!r}"
        )
        from alembic.config import Config
        from alembic import command
        backend_dir = os.path.join(os.path.dirname(__file__), "..")
        cfg = Config(os.path.join(backend_dir, "alembic.ini"))
        cfg.set_main_option("script_location", os.path.join(backend_dir, "migrations"))
        command.upgrade(cfg, "head")
        from fastapi.testclient import TestClient
        from backend.app.main import app
        cls.client_factory = staticmethod(lambda: TestClient(app))
        cls._n = 0

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.test_dir, ignore_errors=True)

    # ------------------------------------------------------------ helpers
    def _org_and_client(self, label):
        type(self)._n += 1
        client = self.client_factory()
        name = f"{label} {self._n} Co"
        r = client.post("/api/auth/signup", json={
            "email": f"intake{self._n}@example.com", "password": "correct-horse-battery-staple",
            "organization_name": name, "display_name": "Intake Test"})
        self.assertIn(r.status_code, (200, 201), r.text)
        from backend.app.repositories import organizations_repo
        org_id = organizations_repo.find_organizations_by_exact_name(name)[0]["id"]
        return org_id, client

    def _case(self, client, spec):
        r = client.post("/api/change-cases", json={
            "name": "case", "trigger_type": "regulatory_restriction",
            "restricted_substance": "PFHxA", "qualification_spec": spec})
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()["id"]

    def _setup(self, label, spec):
        org_id, client = self._org_and_client(label)
        return org_id, client, self._case(client, spec)

    def _upload(self, client, case_id, text, options=None, preview=False, raw=None):
        data = {"options": json.dumps(options)} if options is not None else None
        content = raw if raw is not None else text.encode("utf-8")
        suffix = "/dataset/preview" if preview else "/dataset"
        return client.post(f"/api/change-cases/{case_id}{suffix}",
                           files={"file": ("history.csv", content, "text/csv")}, data=data)

    def _datasets(self, org_id, case_id):
        from backend.app.repositories import qualification_dataset_repo
        return qualification_dataset_repo.list_datasets_for_change_case(org_id, case_id)

    def _experiments(self, org_id, dataset_id):
        from backend.app.repositories import qualification_dataset_repo
        return qualification_dataset_repo.list_experiments_for_dataset(org_id, dataset_id)

    def _assert_rejected_and_nothing_stored(self, org_id, case_id, resp, code):
        self.assertEqual(resp.status_code, 400, resp.text)
        self.assertIn(code, resp.json()["error"])
        self.assertEqual(self._datasets(org_id, case_id), [], "a rejected upload must store no dataset")

    def _paid(self, org_id):
        from backend.app.services import operator_service
        operator_service.activate_subscription(org_id=org_id, plan="pilot")

    def _docx_text(self, content):
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)

    # ------------------------------------------- 1. alternate terminology
    def test_01_alternate_column_terminology_needs_confirmation_then_works(self):
        org_id, client, case_id = self._setup("Alias", SPEC)
        header = ["Crosslinker Ratio", "Cure Temp (C)", "Salt Spray (h)"]
        text = make_csv(header, base_rows())
        self._assert_rejected_and_nothing_stored(org_id, case_id, self._upload(client, case_id, text), "MAPPING_NOT_CONFIRMED")
        pv = self._upload(client, case_id, text, preview=True)
        self.assertEqual(pv.status_code, 200)
        body = pv.json()
        self.assertFalse(body["would_be_accepted"])
        proposed = body["review"]["proposed_mapping"]
        self.assertEqual(proposed["crosslinker_ratio"], "Crosslinker Ratio")
        self.assertEqual(self._datasets(org_id, case_id), [], "preview must persist nothing")
        ok = self._upload(client, case_id, text, options={"column_mapping": proposed})
        self.assertEqual(ok.status_code, 201, ok.text)
        ds = self._datasets(org_id, case_id)[0]
        exps = self._experiments(org_id, ds["id"])
        self.assertEqual(len(exps), 12)
        self.assertEqual(set(json.loads(exps[0]["features_json"])), {"crosslinker_ratio", "cure_temp_c"})
        review = json.loads(ds["review_json"])
        self.assertEqual(review["column_mapping"]["salt_spray_hours"]["method"], "confirmed_mapping")

    # ------------------------------------------- 2. ambiguous terminology
    def test_02_ambiguous_column_terminology_is_never_guessed(self):
        org_id, client, case_id = self._setup("Ambig", SPEC)
        header = ["Crosslinker Ratio", "crosslinker-ratio", "cure_temp_c", "salt_spray_hours"]
        text = make_csv(header, [(a, a, b, c) for a, b, c in base_rows()])
        self._assert_rejected_and_nothing_stored(org_id, case_id, self._upload(client, case_id, text), "AMBIGUOUS_COLUMN_MAPPING")

    # ------------------------------------------- 3. values containing units
    def test_03_values_with_units_ingest_as_canonical_numbers(self):
        org_id, client, case_id = self._setup("Units", SPEC_H)
        text = make_csv(CANON, [(a, b, f"{c} h") for a, b, c in base_rows()])
        r = self._upload(client, case_id, text)
        self.assertEqual(r.status_code, 201, r.text)
        exps = self._experiments(org_id, r.json()["dataset_id"])
        self.assertEqual([e["target_value"] for e in exps], [float(c) for _, _, c in base_rows()])

    # ------------------------------------------- 4. equivalent units
    def test_04_equivalent_units_are_converted_exactly_and_disclosed(self):
        org_id, client, case_id = self._setup("Equiv", SPEC_H)
        rows = [(a, b, c) for a, b, c in base_rows()]
        rows[0] = (rows[0][0], rows[0][1], "24 h")
        rows[1] = (rows[1][0], rows[1][1], "1440 min")
        rows[2] = (rows[2][0], rows[2][1], "1 day")
        r = self._upload(client, case_id, make_csv(CANON, rows), options={"declared_units": {"salt_spray_hours": "h"}})
        self.assertEqual(r.status_code, 201, r.text)
        exps = self._experiments(org_id, r.json()["dataset_id"])
        self.assertEqual([e["target_value"] for e in exps[:3]], [24.0, 24.0, 24.0])
        self.assertIn("UNIT_CONVERTED", {w["code"] for w in r.json()["warnings"]})
        self.assertEqual(r.json()["review"]["units"]["salt_spray_hours"]["rows_converted"], 2)

    # ------------------------------------------- 5. incompatible units
    def test_05_incompatible_units_are_rejected(self):
        org_id, client, case_id = self._setup("Incompat", SPEC_H)
        rows = [(a, b, f"{c} h") for a, b, c in base_rows()]
        rows[3] = (rows[3][0], rows[3][1], "60 %")
        self._assert_rejected_and_nothing_stored(org_id, case_id, self._upload(client, case_id, make_csv(CANON, rows)), "UNIT_INCOMPATIBLE")

    # ------------------------------------------- 6. mixed units
    def test_06_mixed_units_are_rejected_with_instruction(self):
        org_id, client, case_id = self._setup("Mixed", SPEC)
        rows = [(a, b, f"{c} h") for a, b, c in base_rows()[:6]] + [(a, b, f"{2 + i} day") for i, (a, b, c) in enumerate(base_rows()[6:])]
        r = self._upload(client, case_id, make_csv(CANON, rows))
        self._assert_rejected_and_nothing_stored(org_id, case_id, r, "UNIT_MIXED")
        self.assertIn("qualification_spec", r.json()["error"])

    # ------------------------------------------- 7. missing units
    def test_07_missing_units_rejected_when_required_and_resolvable_by_declaration(self):
        org_id, client, case_id = self._setup("Missing", SPEC_H)
        text = make_csv(CANON, base_rows())
        self._assert_rejected_and_nothing_stored(org_id, case_id, self._upload(client, case_id, text), "UNIT_MISSING")
        ok = self._upload(client, case_id, text, options={"declared_units": {"salt_spray_hours": "h"}})
        self.assertEqual(ok.status_code, 201, ok.text)
        review = json.loads(self._datasets(org_id, case_id)[0]["review_json"])
        self.assertEqual(review["units"]["salt_spray_hours"]["declared_source_unit"], "h")

    # ------------------------------------------- 8. different test conditions
    def test_08_different_test_conditions_are_not_merged(self):
        org_id, client, case_id = self._setup("Cond", SPEC)
        rows = [r + (("25" if i < 9 else "80"),) for i, r in enumerate(base_rows())]
        text = make_csv(CANON + ["Test Temp"], rows)
        r = self._upload(client, case_id, text, options={"condition_columns": ["Test Temp"]})
        self._assert_rejected_and_nothing_stored(org_id, case_id, r, "CONDITIONS_DIFFER")
        ok = self._upload(client, case_id, text, options={"condition_columns": ["Test Temp"],
                                                          "reference_conditions": {"Test Temp": "25"}})
        self.assertEqual(ok.status_code, 201, ok.text)
        self.assertEqual(ok.json()["rows_ingested"], 9)
        self.assertEqual(len(self._experiments(org_id, ok.json()["dataset_id"])), 9)
        self.assertEqual(ok.json()["review"]["conditions"]["rows_excluded_other_conditions"], 3)

    # ------------------------------------------- 9. conflicting historical targets
    def test_09_conflicting_targets_rejected_until_explicitly_excluded_with_reason(self):
        org_id, client, case_id = self._setup("Conflict", SPEC)
        rows = base_rows(10) + [(0.2, 200, 100), (0.2, 200, 900), (0.2, 200, 5)]
        text = make_csv(CANON, rows)
        self._assert_rejected_and_nothing_stored(org_id, case_id, self._upload(client, case_id, text), "CONFLICTING_OBSERVATIONS")
        no_reason = self._upload(client, case_id, text, options={"exclude_rows": [12, 13]})
        self._assert_rejected_and_nothing_stored(org_id, case_id, no_reason, "INVALID_OPTIONS")
        ok = self._upload(client, case_id, text, options={"exclude_rows": [12, 13], "exclusion_reason": "typos per customer"})
        self.assertEqual(ok.status_code, 201, ok.text)
        review = json.loads(self._datasets(org_id, case_id)[0]["review_json"])
        self.assertEqual(review["exclusions"]["operator_excluded_rows"], [12, 13])
        self.assertEqual(review["exclusions"]["reason"], "typos per customer")

    # ------------------------------------------- 10. duplicate records
    def test_10_duplicate_records_are_counted_and_not_hidden(self):
        org_id, client, case_id = self._setup("Dupes", SPEC)
        rows = base_rows(12) + base_rows(3)  # 3 exact repeats
        r = self._upload(client, case_id, make_csv(CANON, rows))
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(r.json()["duplicate_rows_found"], 3)
        self.assertEqual(r.json()["rows_ingested"], 15)
        self.assertNotIn("CONFLICTING_OBSERVATIONS", {w["code"] for w in r.json()["review"]["flags"]})

    # ------------------------------------------- 11. messy spreadsheet-style data
    def test_11_messy_spreadsheet_export_is_handled(self):
        org_id, client, case_id = self._setup("Messy", SPEC)
        body = "\n".join(f" {a} , {b} , {c} ,," for a, b, c in base_rows(6)) + "\n,,,,\n,,,,\n" \
            + "\n".join(f" {a} , {b} , {c} ,," for a, b, c in base_rows(12)[6:]) + "\n"
        text = " crosslinker_ratio , cure_temp_c , salt_spray_hours ,,\n" + body
        r = self._upload(client, case_id, "", raw=text.encode("utf-8-sig"))  # with a BOM, like Excel
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(r.json()["rows_ingested"], 12)
        self.assertEqual(r.json()["review"]["counts"]["blank_rows_ignored"], 2)
        bad = self._upload(client, case_id, "Salt spray results Q3,,\n" + make_csv(CANON, base_rows()))
        self.assertEqual(bad.status_code, 400)
        self.assertIn("MISSING_REQUIRED_COLUMN", bad.json()["error"])

    # ------------------------------------------- 12. irrelevant columns
    def test_12_irrelevant_columns_are_ignored_safely_and_disclosed(self):
        org_id, client, case_id = self._setup("Extra", SPEC)
        rows = [r + ("J. Smith", "B117") for r in base_rows()]
        r = self._upload(client, case_id, make_csv(CANON + ["operator", "Test Method"], rows))
        self.assertEqual(r.status_code, 201, r.text)
        codes = {w["code"] for w in r.json()["warnings"]}
        self.assertIn("COLUMNS_IGNORED", codes)
        self.assertIn("POSSIBLE_UNIT_OR_CONDITION_COLUMN_IGNORED", codes)  # Test Method must not vanish silently
        exps = self._experiments(org_id, r.json()["dataset_id"])
        self.assertEqual(set(json.loads(exps[0]["features_json"])), {"crosslinker_ratio", "cure_temp_c"})

    # ------------------------------------------- 13. missing values
    def test_13_missing_values_are_skipped_counted_and_reported(self):
        org_id, client, case_id = self._setup("Gaps", SPEC)
        rows = base_rows(14)
        rows[2] = (rows[2][0], rows[2][1], "")
        rows[7] = (rows[7][0], "", rows[7][2])
        r = self._upload(client, case_id, make_csv(CANON, rows))
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(r.json()["rows_ingested"], 12)
        self.assertEqual(r.json()["rows_skipped"], 2)
        self.assertEqual(sorted(e["row"] for e in r.json()["errors"]), [3, 8])

    # ------------------------------------------- 15. non-finite values
    def test_15_non_finite_values_never_reach_the_model(self):
        org_id, client, case_id = self._setup("Nan", SPEC)
        rows = base_rows(14)
        rows[1] = (rows[1][0], rows[1][1], "nan")
        rows[4] = (rows[4][0], "Infinity", rows[4][2])
        r = self._upload(client, case_id, make_csv(CANON, rows))
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(r.json()["rows_ingested"], 12)
        for e in self._experiments(org_id, r.json()["dataset_id"]):
            self.assertTrue(abs(e["target_value"]) < float("inf") and e["target_value"] == e["target_value"])
        all_bad = self._upload(client, case_id, make_csv(CANON, [(a, b, "nan") for a, b, _ in base_rows()]))
        self.assertEqual(all_bad.status_code, 400)

    # ------------------------------------------- 14. out-of-domain candidates + full path
    def test_14_out_of_domain_candidate_end_to_end_with_review_in_report(self):
        org_id, client, case_id = self._setup("Domain", SPEC_H)
        rows = [(a, b, f"{c} h") for a, b, c in base_rows(15)]
        text = make_csv(CANON, rows)
        raw = text.encode("utf-8")
        up = self._upload(client, case_id, "", raw=raw, options={"declared_units": {}})
        self.assertEqual(up.status_code, 201, up.text)
        client.post(f"/api/change-cases/{case_id}/candidates", json={"name": "Inside", "features": {"crosslinker_ratio": 0.12, "cure_temp_c": 160}})
        client.post(f"/api/change-cases/{case_id}/candidates", json={"name": "Outside", "features": {"crosslinker_ratio": 0.5, "cure_temp_c": 300}})
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 400)  # trial org
        self._paid(org_id)
        rk = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(rk.status_code, 200, rk.text)
        by = {x["candidate_name"]: x for x in rk.json()}
        self.assertEqual(by["Outside"]["decision_support"]["status"], "requires_validation")
        self.assertEqual(by["Inside"]["decision_support"]["status"], "evidence_supported")
        text_doc = self._docx_text(client.post(f"/api/change-cases/{case_id}/report").content)
        self.assertIn("Evidence intake review", text_doc)
        self.assertIn(hashlib.sha256(raw).hexdigest(), text_doc)
        self.assertIn("Requires validation", text_doc)
        self.assertIn("No unit is recorded for: crosslinker_ratio, cure_temp_c", text_doc)
        self.assertNotIn("Units of measure and test conditions are not recorded or verified", text_doc)

    # ------------------------------------------- report honesty for legacy / unreviewed paths
    def test_report_for_plain_canonical_upload_states_units_are_not_recorded(self):
        org_id, client, case_id = self._setup("Plain", SPEC)
        self.assertEqual(self._upload(client, case_id, make_csv(CANON, base_rows(15))).status_code, 201)
        client.post(f"/api/change-cases/{case_id}/candidates", json={"name": "A", "features": {"crosslinker_ratio": 0.12, "cure_temp_c": 160}})
        self._paid(org_id)
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 200)
        t = self._docx_text(client.post(f"/api/change-cases/{case_id}/report").content)
        self.assertIn("Evidence intake review", t)
        self.assertIn("UNITS_NOT_DECLARED", t)
        self.assertIn("No unit is recorded for: crosslinker_ratio, cure_temp_c, salt_spray_hours", t)
        self.assertIn("No test-condition columns were declared at intake", t)

    def test_report_shows_mapping_conversions_conditions_and_exclusions(self):
        org_id, client, case_id = self._setup("Full", SPEC_H)
        header = ["Crosslinker Ratio", "Cure Temp (C)", "Salt Spray (h)", "Test Temp", "source_document", "source_page"]
        rows = []
        for i, (a, b, c) in enumerate(base_rows(15)):
            t = "25" if i < 13 else "80"
            cell = "1 day" if i == 0 else f"{c} h"
            rows.append((a, b, cell, t, "Report 12", 4))
        rows += [(0.3, 300, "100 h", "25", "Report 9", 1), (0.3, 300, "900 h", "25", "Report 9", 2)]
        opts = {"column_mapping": {"crosslinker_ratio": "Crosslinker Ratio", "cure_temp_c": "Cure Temp (C)", "salt_spray_hours": "Salt Spray (h)"},
                "condition_columns": ["Test Temp"], "reference_conditions": {"Test Temp": "25"},
                "exclude_rows": [17], "exclusion_reason": "Report 9 p.2 superseded by customer"}
        up = self._upload(client, case_id, make_csv(header, rows), options=opts)
        self.assertEqual(up.status_code, 201, up.text)
        client.post(f"/api/change-cases/{case_id}/candidates", json={"name": "A", "features": {"crosslinker_ratio": 0.12, "cure_temp_c": 160}})
        self._paid(org_id)
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 200)
        t = self._docx_text(client.post(f"/api/change-cases/{case_id}/report").content)
        for needle in ("Reviewer-confirmed mapping", "Salt Spray (h)", "1 row(s) day -> h", "Test Temp",
                       "Reference condition set used: Test Temp=25", "Report 9 p.2 superseded by customer",
                       "source_document, source_page", "Rows excluded (other test conditions)"):
            self.assertIn(needle, t, needle)

    # ------------------------------------------- route-level safety
    def test_preview_and_upload_are_tenant_scoped(self):
        _, client, case_id = self._setup("TenantA", SPEC)
        _, other = self._org_and_client("TenantB")
        text = make_csv(CANON, base_rows())
        self.assertEqual(self._upload(other, case_id, text, preview=True).status_code, 404)
        self.assertEqual(self._upload(other, case_id, text).status_code, 404)

    def test_invalid_options_json_is_a_clean_400(self):
        org_id, client, case_id = self._setup("BadJson", SPEC)
        r = client.post(f"/api/change-cases/{case_id}/dataset",
                        files={"file": ("h.csv", make_csv(CANON, base_rows()).encode(), "text/csv")},
                        data={"options": "{not json"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self._datasets(org_id, case_id), [])

    def test_invalid_spec_units_rejected_at_case_creation(self):
        _, client = self._org_and_client("BadSpec")
        for bad in ({"units": ["h"]}, {"units": {"not_a_column": "h"}}, {"units": {"salt_spray_hours": ""}}):
            r = client.post("/api/change-cases", json={
                "name": "c", "trigger_type": "regulatory_restriction", "restricted_substance": "x",
                "qualification_spec": {**SPEC, **bad}})
            self.assertEqual(r.status_code, 400, bad)

    def test_upload_without_options_behaves_like_before_for_clean_canonical_data(self):
        org_id, client, case_id = self._setup("Legacy", SPEC)
        r = self._upload(client, case_id, make_csv(CANON, base_rows(15)))
        self.assertEqual(r.status_code, 201, r.text)
        for key in ("dataset_id", "rows_ingested", "rows_skipped", "errors", "data_quality_status",
                    "duplicate_rows_found", "constant_columns"):
            self.assertIn(key, r.json())
        self.assertEqual(r.json()["rows_ingested"], 15)


if __name__ == "__main__":
    unittest.main()
