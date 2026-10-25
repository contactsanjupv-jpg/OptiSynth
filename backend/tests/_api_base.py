"""
Shared base for API tests that need their own isolated SQLite database.
Same rebind-and-assert isolation as test_api.py / test_change_case_api.py, so
every module using it is safe alongside any other module in one combined run.
"""
import io
import json
import os
import re
import shutil
import unittest

MIME = {
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
}

# A diagnostic can only be ranked when EVERY modelled field and the target has an explicit canonical unit, so the
# default fixture declares all three. (crosslinker_ratio is a dimensionless fraction.)
SPEC = {"feature_columns": ["crosslinker_ratio", "cure_temp_c"], "target_metric": "salt_spray_hours",
        "target_value": 300.0, "direction": "maximize",
        "units": {"crosslinker_ratio": "fraction", "cure_temp_c": "degC", "salt_spray_hours": "h"}}
# The same case with NO units declared: cannot be ranked (see test_units_before_ranking.py).
SPEC_NO_UNITS = {k: v for k, v in SPEC.items() if k != "units"}
# Partly declared specs: kept because the unit-conversion tests are about exactly that; they cannot be ranked.
SPEC_H = {**SPEC_NO_UNITS, "units": {"salt_spray_hours": "h"}}
SPEC_U = {**SPEC_NO_UNITS, "units": {"cure_temp_c": "degC", "salt_spray_hours": "h"}}

CONFIRM_ALIASES = {"column_mapping": {
    "crosslinker_ratio": "Crosslinker Ratio", "cure_temp_c": "Cure Temp (C)", "salt_spray_hours": "Salt Spray (h)"}}


_CANDIDATE_URL = re.compile(r"/api/change-cases/(\d+)/candidates(?:/\d+)?$")
_DATASET_URL = re.compile(r"/api/change-cases/(\d+)/dataset(?:/preview)?$")


def unit_aware_client(app):
    """A TestClient for tests that are not ABOUT units but need a rankable case.

    Ranking requires every modelled field to have a declared canonical unit, and a case that declares one
    requires each candidate and each uploaded file to state it. So that the many tests that merely add a
    candidate or upload a default file need not repeat that, this client states the CASE'S OWN declared unit on
    their behalf -- identity, never a guess or a conversion. It steps aside whenever a test says something
    itself: a candidate request that has a `units` key (even an empty one), or an upload whose options have a
    `declared_units` key (even an empty one), is sent exactly as written. The product code has no such shortcut."""
    from fastapi.testclient import TestClient

    class UnitAwareTestClient(TestClient):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._case_units = {}

        def _units_of(self, case_id):
            if case_id not in self._case_units:
                response = TestClient.request(self, "GET", f"/api/change-cases/{case_id}")
                units = {}
                if response.status_code == 200:
                    units = (json.loads(response.json().get("qualification_spec_json") or "{}").get("units") or {})
                self._case_units[case_id] = units
            return self._case_units[case_id]

        def request(self, method, url, **kwargs):
            path = str(url).split("?")[0]
            if method.upper() in ("POST", "PATCH"):
                m = _CANDIDATE_URL.search(path)
                body = kwargs.get("json")
                if m and isinstance(body, dict) and isinstance(body.get("features"), dict) and "units" not in body:
                    declared = self._units_of(int(m.group(1)))
                    stated = {f: declared[f] for f in body["features"] if f in declared}
                    if stated:
                        kwargs["json"] = {**body, "units": stated}
                m = _DATASET_URL.search(path)
                if m and "files" in kwargs:
                    data = dict(kwargs.get("data") or {})
                    try:
                        options = json.loads(data["options"]) if data.get("options") else {}
                    except ValueError:
                        options = None  # invalid options JSON: send it exactly as written (a test of that error)
                    if isinstance(options, dict) and "declared_units" not in options:  # a test that sets declared_units (even {}) is left alone
                        declared = self._units_of(int(m.group(1)))
                        if declared:
                            kwargs["data"] = {**data, "options": json.dumps({**options, "declared_units": declared})}
            return super().request(method, url, **kwargs)

    return UnitAwareTestClient(app)


def persist_legacy_ranking(organization_id, change_case_id):
    """Stores one prediction per candidate directly through the repository, exactly as a ranking that was saved
    BEFORE units were required would look. The API now refuses to rank a diagnostic whose fields have no declared
    unit, so tests of how the report treats such a legacy diagnostic cannot get its predictions from /rank."""
    from backend.app.repositories import candidate_repo, qualification_dataset_repo
    dataset = qualification_dataset_repo.get_latest_dataset_for_change_case(organization_id, change_case_id)
    candidates = candidate_repo.list_candidates_for_change_case(organization_id, change_case_id)
    candidate_repo.persist_ranking(organization_id, dataset["id"], [
        {"candidate_id": c["id"], "model_version": "gp-rank-v1", "predicted_value": 400.0 + i,
         "predicted_probability": round(0.9 - 0.1 * i, 2), "uncertainty_std": 5.0,
         "experiment_description": "Run a physical qualification test.", "priority": 1,
         "inputs_version": int(c.get("inputs_version") or 0)}
        for i, c in enumerate(candidates)
    ])


class IsolatedApiTestCase(unittest.TestCase):
    dir_name = "_test_env_iso"

    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.path.dirname(__file__), cls.dir_name)
        os.makedirs(cls.test_dir, exist_ok=True)
        os.environ["SECRET_KEY"] = "test-secret-key-not-for-production"
        os.environ["PASSWORD_PEPPER"] = "test-password-pepper-not-for-production"
        os.environ["DATABASE_URL"] = f"sqlite:///{cls.test_dir}/test.db"
        if os.path.exists(f"{cls.test_dir}/test.db"):
            os.remove(f"{cls.test_dir}/test.db")
        from backend.app.config.settings import settings
        from backend.app.config import database
        settings.DATABASE_URL = os.environ["DATABASE_URL"]
        database.engine = database._make_engine()
        assert str(database.engine.url).endswith(f"{cls.test_dir}/test.db"), (
            f"Test isolation failure: engine bound to {database.engine.url!r}"
        )
        from alembic.config import Config
        from alembic import command
        backend_dir = os.path.join(os.path.dirname(__file__), "..")
        cfg = Config(os.path.join(backend_dir, "alembic.ini"))
        cfg.set_main_option("script_location", os.path.join(backend_dir, "migrations"))
        command.upgrade(cfg, "head")
        from backend.app.main import app
        cls.client_factory = staticmethod(lambda: unit_aware_client(app))
        cls._n = 0

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.test_dir, ignore_errors=True)

    # ------------------------------------------------------------ helpers
    def org_and_client(self, label):
        type(self)._n += 1
        client = self.client_factory()
        name = f"{label} {self._n} Co"
        r = client.post("/api/auth/signup", json={
            "email": f"{self.dir_name}{self._n}@example.com", "password": "correct-horse-battery-staple",
            "organization_name": name, "display_name": "Test User"})
        assert r.status_code in (200, 201), r.text
        from backend.app.repositories import organizations_repo
        org_id = organizations_repo.find_organizations_by_exact_name(name)[0]["id"]
        self.__class__._org_of_client = {**getattr(self.__class__, "_org_of_client", {}), id(client): org_id}
        return org_id, client

    def make_legacy_case(self, org_id, spec, name="case"):
        """Stores a diagnostic DIRECTLY through the repository, bypassing the API's validation: the only way a
        diagnostic without a full set of declared units can exist now (the API refuses to create one). It is exactly
        what an older diagnostic created before units were mandatory looks like."""
        from sqlalchemy import text
        from backend.app.config.database import db_connection
        from backend.app.repositories import change_case_repo
        with db_connection() as conn:
            user_id = conn.execute(text("SELECT user_id FROM memberships WHERE organization_id = :o ORDER BY id LIMIT 1"),
                                   {"o": org_id}).scalar()
        return change_case_repo.create_change_case(org_id, user_id, name, "regulatory_restriction", "PFHxA",
                                                   json.dumps(spec))

    def make_case(self, client, spec):
        from backend.app.services.change_case_rules import units_declaration_problems
        if units_declaration_problems(spec):
            # A spec without every unit can no longer be created through the API; the tests that need one are
            # about legacy diagnostics (or about unit handling that does not depend on creation).
            return self.make_legacy_case(self._org_of_client[id(client)], spec)
        r = client.post("/api/change-cases", json={
            "name": "case", "trigger_type": "regulatory_restriction",
            "restricted_substance": "PFHxA", "qualification_spec": spec})
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()["id"]

    def setup_case(self, label, spec=SPEC):
        org_id, client = self.org_and_client(label)
        return org_id, client, self.make_case(client, spec)

    def post_file(self, client, case_id, filename, content, options=None, route="dataset"):
        ext = os.path.splitext(filename)[1].lower()
        data = {"options": json.dumps(options)} if options is not None else None
        return client.post(f"/api/change-cases/{case_id}/{route}",
                           files={"file": (filename, content, MIME.get(ext, "application/octet-stream"))}, data=data)

    def upload(self, client, case_id, filename, content, options=None):
        return self.post_file(client, case_id, filename, content, options, "dataset")

    def preview(self, client, case_id, filename, content, options=None):
        return self.post_file(client, case_id, filename, content, options, "dataset/preview")

    def extract(self, client, case_id, filename, content):
        return self.post_file(client, case_id, filename, content, None, "evidence/extract")

    def datasets(self, org_id, case_id):
        from backend.app.repositories import qualification_dataset_repo
        return qualification_dataset_repo.list_datasets_for_change_case(org_id, case_id)

    def experiments(self, org_id, dataset_id):
        from backend.app.repositories import qualification_dataset_repo
        return qualification_dataset_repo.list_experiments_for_dataset(org_id, dataset_id)

    def paid(self, org_id):
        from backend.app.services import operator_service
        operator_service.activate_subscription(org_id=org_id, plan="pilot")

    def assert_rejected_nothing_stored(self, org_id, case_id, resp, code):
        self.assertEqual(resp.status_code, 400, resp.text)
        self.assertIn(code, resp.json()["error"])
        self.assertEqual(self.datasets(org_id, case_id), [], "a rejected upload must store nothing")

    def docx_text(self, content):
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
