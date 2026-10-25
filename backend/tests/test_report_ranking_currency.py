"""
A generated qualification report must never present a comparison of an EARLIER candidate set as the current
ranking, and must never set a recorded physical result beside a model estimate made after it.

THE DEFECT. generate_report() kept only candidates that have a prediction for their current inputs, then
ranked those itself. It never consulted the candidate listing's "ranking_outdated" state. So after the
candidate set changed -- a candidate added, removed, restored, or any candidate's inputs edited -- the report
still printed a ranked table whose positions came from a run over a DIFFERENT set (a new candidate silently
missing, a removed one silently absent, positions renumbered among whoever was left) while the candidate
listing correctly withheld those positions. The outcomes table had a related fault: it paired each recorded
physical result with the candidate's LATEST estimate, even when that estimate was produced after the result
was recorded -- implying the model had predicted a test that had already happened.

THE BEHAVIOUR UNDER TEST. Exactly the existing refusal convention (mixed dataset versions are already refused
with "Re-run the ranking ... before generating a report"): when the last ranking no longer matches the current
candidates, the report is refused with an actionable error, and no partial or stale ranking is produced.
Re-running the ranking makes it possible again. Outcomes are compared only with an estimate that existed
before the result was recorded.

Synthetic data: these tests prove the logic and that the report agrees with the candidate listing; they say
nothing about the scientific validity of the model.

Run:  python3 -m unittest backend.tests.test_report_ranking_currency -v
"""
import io

from backend.tests import evidence_fixtures as F
from backend.tests._api_base import IsolatedApiTestCase

HISTORY = F.base_rows(14)  # crosslinker 0.080-0.145, cure temp 140-179 C
IN_A = {"crosslinker_ratio": 0.12, "cure_temp_c": 158}
IN_B = {"crosslinker_ratio": 0.10, "cure_temp_c": 150}
IN_C = {"crosslinker_ratio": 0.13, "cure_temp_c": 165}
IN_A_EDITED = {"crosslinker_ratio": 0.09, "cure_temp_c": 146}
FAR = {"crosslinker_ratio": 0.40, "cure_temp_c": 300}
FAR_BELOW = {"crosslinker_ratio": 0.00, "cure_temp_c": 60}
NOT_RANKED = "not_ranked_insufficient_evidence"


class TestReportRankingCurrency(IsolatedApiTestCase):
    dir_name = "_test_env_report_currency"

    # ---------------------------------------------------------------- helpers
    def _case(self, label, candidates=()):
        org_id, client, case_id = self.setup_case(label)
        self.paid(org_id)
        up = self.upload(client, case_id, "history.csv", F.csv_bytes(F.CANON_HEADERS, HISTORY))
        self.assertEqual(up.status_code, 201, up.text)
        ids = {name: self._add(client, case_id, name, features) for name, features in candidates}
        return org_id, client, case_id, ids

    def _add(self, client, case_id, name, features):
        r = client.post(f"/api/change-cases/{case_id}/candidates", json={"name": name, "features": features})
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()["id"]

    def _rank(self, client, case_id):
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200, r.text)
        return {x["candidate_name"]: x for x in r.json()}

    def _list(self, client, case_id):
        r = client.get(f"/api/change-cases/{case_id}/candidates")
        self.assertEqual(r.status_code, 200, r.text)
        return {c["candidate_name"]: c for c in r.json()}

    def _report_raw(self, client, case_id):
        return client.post(f"/api/change-cases/{case_id}/report")

    def _report(self, client, case_id):
        from docx import Document
        r = self._report_raw(client, case_id)
        self.assertEqual(r.status_code, 200, r.text[:300])
        doc = Document(io.BytesIO(r.content))
        paragraphs = [p.text for p in doc.paragraphs]
        tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
        return paragraphs, tables

    @staticmethod
    def _table(tables, first_header):
        found = [t for t in tables if t[0][0] == first_header]
        return found[0] if found else None

    @staticmethod
    def _unranked_table(tables):
        """The "Not ranked: outside the historical evidence" table (its headers say "extrapolation")."""
        found = [t for t in tables if any("extrapolation" in h for h in t[0])]
        return found[0] if found else None

    def _assert_refused(self, client, case_id, names=()):
        r = self._report_raw(client, case_id)
        self.assertEqual(r.status_code, 400, "the report must be refused, not generated with a stale ranking")
        self.assertIn("application/json", r.headers["content-type"], "no document may be produced")
        message = r.json()["error"]
        self.assertIn("out of date", message)
        self.assertIn("Re-run ranking", message)
        for name in names:
            self.assertIn(name, message)
        return message

    def _assert_report_and_listing_agree(self, client, case_id, refused):
        """The report is refused exactly when the candidate listing says the ranking is out of date."""
        listing = self._list(client, case_id)
        outdated = any(c.get("ranking_outdated") or c.get("prediction_inputs_stale") for c in listing.values())
        self.assertEqual(outdated, refused, f"listing flags disagree with the expected state: {listing}")
        if refused:
            self._assert_refused(client, case_id)
        else:
            self.assertEqual(self._report_raw(client, case_id).status_code, 200)

    # ---------------------------------------------------------------- the report right after ranking
    def test_report_right_after_ranking_shows_the_current_ranking(self):
        _, client, case_id, _ = self._case("fresh", [("InA", IN_A), ("InB", IN_B), ("Far", FAR)])
        by = self._rank(client, case_id)
        paragraphs, tables = self._report(client, case_id)
        ranked = self._table(tables, "Rank")
        names = [row[1] for row in ranked[1:]]
        self.assertEqual(names, sorted((n for n in by if by[n]["rank"] is not None), key=lambda n: by[n]["rank"]))
        self.assertEqual([row[0] for row in ranked[1:]], [str(by[n]["rank"]) for n in names])
        self.assertNotIn("Far", names, "the unsupported candidate stays unranked")
        self.assertEqual([row[0] for row in self._unranked_table(tables)[1:]], ["Far"])
        self.assertFalse(any("out of date" in p.lower() for p in paragraphs))

    # ---------------------------------------------------------------- candidate set changed
    def test_adding_a_candidate_after_ranking_refuses_the_report_instead_of_printing_old_positions(self):
        _, client, case_id, _ = self._case("added", [("InA", IN_A), ("InB", IN_B)])
        self._rank(client, case_id)
        self.assertEqual(self._report_raw(client, case_id).status_code, 200)
        self._add(client, case_id, "InC", IN_C)
        listing = self._list(client, case_id)
        self.assertTrue(all(c["rank"] is None and c["ranking_outdated"] for n, c in listing.items() if n != "InC"),
                        "the listing already withholds positions: the report must agree")
        self._assert_refused(client, case_id)

    def test_removing_and_restoring_follow_the_listings_ranking_state(self):
        _, client, case_id, ids = self._case("setchange", [("InA", IN_A), ("InB", IN_B)])
        base = f"/api/change-cases/{case_id}/candidates"
        self._rank(client, case_id)
        self._assert_report_and_listing_agree(client, case_id, refused=False)

        self.assertEqual(client.post(f"{base}/{ids['InA']}/archive").status_code, 200)
        self._assert_report_and_listing_agree(client, case_id, refused=True)     # a ranked candidate was removed
        self._rank(client, case_id)
        self._assert_report_and_listing_agree(client, case_id, refused=False)
        _, tables = self._report(client, case_id)
        self.assertEqual([row[1] for row in self._table(tables, "Rank")[1:]], ["InB"])

        self.assertEqual(client.post(f"{base}/{ids['InA']}/restore").status_code, 200)
        self._assert_report_and_listing_agree(client, case_id, refused=True)     # a candidate was restored
        self._rank(client, case_id)
        self._assert_report_and_listing_agree(client, case_id, refused=False)
        _, tables = self._report(client, case_id)
        self.assertEqual(sorted(row[1] for row in self._table(tables, "Rank")[1:]), ["InA", "InB"])

    def test_a_candidate_added_and_removed_before_it_was_ever_ranked_leaves_the_ranking_current(self):
        """The established rule: removing a candidate that was never ranked does not change what was ranked."""
        _, client, case_id, _ = self._case("transient", [("InA", IN_A), ("InB", IN_B)])
        self._rank(client, case_id)
        extra = self._add(client, case_id, "InC", IN_C)
        self._assert_report_and_listing_agree(client, case_id, refused=True)
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/candidates/{extra}/archive").status_code, 200)
        self._assert_report_and_listing_agree(client, case_id, refused=False)

    def test_rerunning_the_ranking_after_an_addition_restores_a_valid_current_report(self):
        _, client, case_id, _ = self._case("rerun", [("InA", IN_A), ("InB", IN_B)])
        self._rank(client, case_id)
        self._add(client, case_id, "InC", IN_C)
        self._assert_refused(client, case_id)
        by = self._rank(client, case_id)
        _, tables = self._report(client, case_id)
        ranked = self._table(tables, "Rank")
        self.assertEqual({row[1] for row in ranked[1:]}, {"InA", "InB", "InC"})
        self.assertEqual({row[1]: row[0] for row in ranked[1:]}, {n: str(by[n]["rank"]) for n in by})

    # ---------------------------------------------------------------- edited inputs
    def test_editing_inputs_after_ranking_never_lets_the_old_prediction_stand_as_current(self):
        _, client, case_id, ids = self._case("edited", [("InA", IN_A), ("InB", IN_B)])
        before = self._rank(client, case_id)
        r = client.patch(f"/api/change-cases/{case_id}/candidates/{ids['InA']}", json={"features": IN_A_EDITED})
        self.assertEqual(r.status_code, 200, r.text)
        message = self._assert_refused(client, case_id, names=["InA"])
        self.assertIn("Input values were changed", message)
        by = self._rank(client, case_id)
        self.assertNotEqual(by["InA"]["predicted_value"], before["InA"]["predicted_value"],
                            "the edited inputs produce a different prediction (precondition of this test)")
        _, tables = self._report(client, case_id)
        ranked = self._table(tables, "Rank")
        row = next(row for row in ranked[1:] if row[1] == "InA")
        shown = row[2]
        self.assertIn(str(round(by["InA"]["predicted_value"]))[:3], shown, "the new prediction is what is shown")
        self.assertNotIn(str(round(before["InA"]["predicted_value"])), shown)

    def test_editing_the_only_candidate_is_refused_rather_than_reported_as_not_yet_ranked(self):
        _, client, case_id, ids = self._case("onlyone", [("Solo", IN_A)])
        self._rank(client, case_id)
        client.patch(f"/api/change-cases/{case_id}/candidates/{ids['Solo']}", json={"features": IN_A_EDITED})
        self._assert_refused(client, case_id, names=["Solo"])
        self._rank(client, case_id)
        self.assertEqual(self._report_raw(client, case_id).status_code, 200)

    # ---------------------------------------------------------------- domain-aware behaviour is preserved
    def test_entirely_out_of_domain_candidates_stay_unranked_through_the_refuse_and_rerun_cycle(self):
        _, client, case_id, _ = self._case("alloutside", [("FarA", FAR), ("FarB", FAR_BELOW)])
        self._rank(client, case_id)
        paragraphs, tables = self._report(client, case_id)
        self.assertIn("No candidate could be ranked", " ".join(paragraphs))
        self.assertIsNone(self._table(tables, "Rank"))
        self._add(client, case_id, "FarC", {"crosslinker_ratio": 0.90, "cure_temp_c": 400})
        self._assert_refused(client, case_id)
        by = self._rank(client, case_id)
        self.assertTrue(all(r["rank"] is None and r["ranking_status"] == NOT_RANKED for r in by.values()))
        paragraphs, tables = self._report(client, case_id)
        self.assertIn("No candidate could be ranked", " ".join(paragraphs))
        self.assertIsNone(self._table(tables, "Rank"), "no position may be invented for unsupported candidates")
        unranked = self._unranked_table(tables)
        self.assertEqual({row[0] for row in unranked[1:]}, {"FarA", "FarB", "FarC"})

    def test_supported_candidates_keep_the_evidence_aware_order_and_the_unsupported_one_stays_unranked(self):
        _, client, case_id, _ = self._case("mixed", [("InA", IN_A), ("InB", IN_B), ("Far", FAR)])
        self._rank(client, case_id)
        self._add(client, case_id, "InC", IN_C)
        self._assert_refused(client, case_id)
        by = self._rank(client, case_id)
        paragraphs, tables = self._report(client, case_id)
        ranked = self._table(tables, "Rank")
        names = [row[1] for row in ranked[1:]]
        self.assertEqual(names, sorted((n for n in by if by[n]["rank"] is not None), key=lambda n: by[n]["rank"]))
        self.assertNotIn("Far", names)
        self.assertEqual([row[0] for row in self._unranked_table(tables)[1:]], ["Far"])
        self.assertIn("Not ranked: outside the historical evidence", paragraphs)

    # ---------------------------------------------------------------- recorded outcomes
    def _record_outcome(self, client, case_id, candidate_id, passed=True):
        r = client.post(f"/api/change-cases/{case_id}/candidates/{candidate_id}/outcome",
                        json={"actual_result": "Passed salt spray at 540 h", "passed_spec": passed})
        self.assertEqual(r.status_code, 201, r.text)

    @staticmethod
    def _outcomes_table(tables):
        found = [t for t in tables if t[0][0] == "Candidate" and "Recorded result" in t[0]]
        return found[0] if found else None

    def test_an_outcome_is_compared_only_with_an_estimate_that_existed_before_it_was_recorded(self):
        org_id, client, case_id, ids = self._case("outcome_order", [("InA", IN_A), ("InB", IN_B)])
        self._rank(client, case_id)
        self._record_outcome(client, case_id, ids["InA"])
        self._rank(client, case_id)  # a NEWER estimate for InA now exists
        from backend.app.repositories import candidate_repo
        history = candidate_repo.list_predictions_for_candidate(org_id, ids["InA"])
        self.assertEqual(len(history), 2)
        _, tables = self._report(client, case_id)
        outcomes = self._outcomes_table(tables)
        self.assertIn("before the result was recorded", outcomes[0][1])
        shown = outcomes[1][1]
        self.assertIn(history[0]["created_at"], shown, "the estimate that existed when the result was recorded")
        self.assertNotIn(history[1]["created_at"], shown, "the newer estimate must not be paired with the result")
        self.assertLess(history[0]["created_at"], outcomes[1][4], "the estimate is older than the recorded result")
        self.assertGreater(history[1]["created_at"], outcomes[1][4])

    def test_an_outcome_recorded_before_any_estimate_is_not_compared_with_one(self):
        _, client, case_id, ids = self._case("outcome_first", [("InA", IN_A)])
        self._record_outcome(client, case_id, ids["InA"])  # the API allows it even though nothing was ranked yet
        self._rank(client, case_id)
        _, tables = self._report(client, case_id)
        cell = self._outcomes_table(tables)[1][1]
        self.assertIn("Not compared", cell)
        self.assertIn("no model estimate", cell)
        self.assertNotIn("%", cell)

    def test_the_recommendation_does_not_ask_for_a_test_that_has_already_been_recorded_without_saying_so(self):
        _, client, case_id, ids = self._case("recommend", [("InA", IN_A), ("InB", IN_B)])
        self._rank(client, case_id)
        self._record_outcome(client, case_id, ids["InA"])
        self._rank(client, case_id)
        paragraphs, _ = self._report(client, case_id)
        rec = {p.split(": ", 1)[0]: p for p in paragraphs
               if p.startswith(("InA: Run", "InB: Run"))}
        self.assertIn("physical outcome has already been recorded", rec["InA"])
        self.assertNotIn("already been recorded", rec["InB"])
        gaps = [p for p in paragraphs if p.startswith("No physical validation outcome has been recorded for:")]
        self.assertEqual(len(gaps), 1)
        self.assertIn("InB", gaps[0])
        self.assertNotIn("InA", gaps[0], "the evidence-gap list must not claim a recorded outcome is missing")


if __name__ == "__main__":
    import unittest
    unittest.main()
