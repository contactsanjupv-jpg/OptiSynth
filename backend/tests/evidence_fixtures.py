"""
Adversarial evidence fixtures: realistic-looking files from DIFFERENT imaginary
customers, built at test time (no binary files are committed).

    python3 scripts/make_evidence_fixtures.py <folder>    # writes them to disk to try in the UI

Each builder returns file BYTES. Grids are lists of rows of plain values.
The shared case used by most tests: features crosslinker_ratio, cure_temp_c;
target salt_spray_hours (maximize, 300).
"""
import io

CANON_HEADERS = ["crosslinker_ratio", "cure_temp_c", "salt_spray_hours"]


def base_rows(n=12):
    """n DISTINCT numeric rows: (crosslinker ratio, cure temp C, salt spray hours)."""
    return [(round(0.08 + 0.005 * i, 4), 140 + 3 * i, 100 + 40 * i) for i in range(n)]


# ---------------------------------------------------------------------------
# generic builders
# ---------------------------------------------------------------------------
def make_xlsx(sheets, hidden=(), formats=None):
    """sheets: {name: grid}. formats: {(sheet, row, col_1based): number_format}."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, grid in sheets.items():
        ws = wb.create_sheet(name)
        for r in grid:
            ws.append(list(r))
        if name in hidden:
            ws.sheet_state = "hidden"
    for (sheet, row, col), fmt in (formats or {}).items():
        wb[sheet].cell(row=row, column=col).number_format = fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx(blocks):
    """blocks: list of ("p", text) or ("t", grid)."""
    from docx import Document
    doc = Document()
    for kind, content in blocks:
        if kind == "p":
            doc.add_paragraph(content)
        else:
            table = doc.add_table(rows=len(content), cols=max(len(r) for r in content))
            table.style = "Table Grid"
            for i, row in enumerate(content):
                for j, v in enumerate(row):
                    table.cell(i, j).text = "" if v is None else str(v)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(blocks):
    """blocks: list of ("p", text) | ("t", grid) | ("pagebreak", None). Tables are ruled."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []
    for kind, content in blocks:
        if kind == "p":
            story += [Paragraph(content, styles["Normal"]), Spacer(1, 12)]
        elif kind == "pagebreak":
            story.append(PageBreak())
        else:
            t = Table([["" if v is None else str(v) for v in r] for r in content])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
            story += [t, Spacer(1, 12)]
    doc.build(story)
    return buf.getvalue()


def make_image_only_pdf():
    """A page with a drawn rectangle and NO text layer (what a scan looks like to a parser)."""
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(100, 100, 300, 200)
    c.showPage()
    c.save()
    return buf.getvalue()


# ---------------------------------------------------------------------------
# customer scenarios (grids)
# ---------------------------------------------------------------------------
def grid(headers, rows):
    return [list(headers)] + [list(r) for r in rows]


def customer_alt_headers_grid(n=12):
    """Different terminology + irrelevant columns."""
    rows = [(a, b, c, "J. Smith", "B117") for a, b, c in base_rows(n)]
    return grid(["Crosslinker Ratio", "Cure Temp (C)", "Salt Spray (h)", "Operator", "Test Standard"], rows)


def xlsx_title_rows_then_table():
    """Title rows above the header (header is on source row 3)."""
    g = [["Coil coating study 2024 - salt spray"], [None]] + customer_alt_headers_grid()
    return make_xlsx({"Results": g})


def xlsx_two_sheets():
    return make_xlsx({
        "Summary": grid(["Item", "Value"], [("Lab", "North"), ("Year", 2024), ("Panels", 12)]),
        "Data": grid(CANON_HEADERS, base_rows()),
    })


def xlsx_alt_headers():
    return make_xlsx({"Data": customer_alt_headers_grid()})


def xlsx_mixed_units():
    rows = [(a, b, f"{c} h") for a, b, c in base_rows()]
    rows[0] = (rows[0][0], rows[0][1], "24 h")
    rows[1] = (rows[1][0], rows[1][1], "1 day")
    rows[2] = (rows[2][0], rows[2][1], "1440 min")
    rows[3] = (rows[3][0], rows[3][1], "5 %")
    return make_xlsx({"Data": grid(CANON_HEADERS, rows)})


def xlsx_equivalent_units():
    rows = [(a, b, f"{c} h") for a, b, c in base_rows()]
    rows[0] = (rows[0][0], rows[0][1], "24 h")
    rows[1] = (rows[1][0], rows[1][1], "1 day")
    rows[2] = (rows[2][0], rows[2][1], "1440 min")
    return make_xlsx({"Data": grid(CANON_HEADERS, rows)})


def xlsx_missing_units_with_hidden_number_format():
    """Target values are bare numbers but DISPLAY as '24 h' via the number format."""
    g = grid(CANON_HEADERS, base_rows())
    formats = {("Data", r, 3): '0 "h"' for r in range(2, len(g) + 1)}
    return make_xlsx({"Data": g}, formats=formats)


def xlsx_multiple_conditions():
    rows = [r + (("25" if i < 8 else "80"),) for i, r in enumerate(base_rows(14))]
    return make_xlsx({"Data": grid(CANON_HEADERS + ["Test Temp (C)"], rows)})


def xlsx_duplicates():
    rows = base_rows(12) + base_rows(4)
    return make_xlsx({"Data": grid(CANON_HEADERS, rows)})


def xlsx_conflicting():
    rows = base_rows(10) + [(0.2, 200, 100), (0.2, 200, 900), (0.2, 200, 5)]
    return make_xlsx({"Data": grid(CANON_HEADERS, rows)})


def xlsx_missing_target_column():
    return make_xlsx({"Data": grid(CANON_HEADERS[:2], [(a, b) for a, b, _ in base_rows()])})


def xlsx_hidden_sheet_only_data():
    return make_xlsx({"Cover": grid(["Note"], [("see Data",)]), "Data": grid(CANON_HEADERS, base_rows())}, hidden=("Data",))


def docx_customer_report():
    """A lab report: narrative, a conditions table, then the results table."""
    return make_docx([
        ("p", "Salt spray test report 2024-09"),
        ("p", "Table 1. Panel list"),
        ("t", grid(["Panel", "Substrate"], [("P1", "steel"), ("P2", "steel")])),
        ("p", "Table 2. Salt spray results, ASTM B117, 35 C"),
        ("t", grid(["Crosslinker Ratio", "Cure Temp (C)", "Salt Spray (h)"], base_rows())),
    ])


def docx_single_table():
    return make_docx([("p", "Results"), ("t", grid(CANON_HEADERS, base_rows()))])


def docx_conditions_in_table():
    rows = [r + (("35" if i < 6 else "50"),) for i, r in enumerate(base_rows(12))]
    return make_docx([("p", "Results by test temperature"), ("t", grid(CANON_HEADERS + ["Chamber Temp (C)"], rows))])


def docx_no_tables():
    return make_docx([("p", "Narrative only, no tables.")])


def pdf_ruled_table():
    return make_pdf([("p", "Supplier test certificate"), ("t", grid(["Crosslinker Ratio", "Cure Temp (C)", "Salt Spray (h)"], base_rows()))])


def pdf_two_tables():
    return make_pdf([
        ("p", "Page one"),
        ("t", grid(["Item", "Value"], [("Lab", "North"), ("Year", 2024)])),
        ("pagebreak", None),
        ("p", "Page two results"),
        ("t", grid(CANON_HEADERS, base_rows())),
    ])


def pdf_table_continued_over_two_pages():
    rows = base_rows(12)
    return make_pdf([
        ("t", grid(CANON_HEADERS, rows[:6])),
        ("pagebreak", None),
        ("t", grid(CANON_HEADERS, rows[6:])),
    ])


def pdf_scanned():
    return make_image_only_pdf()


def csv_bytes(headers, rows):
    return (",".join(headers) + "\n" + "\n".join(",".join(str(c) for c in r) for r in rows) + "\n").encode("utf-8")


ALL_FIXTURES = {
    "xlsx_alt_headers.xlsx": xlsx_alt_headers,
    "xlsx_title_rows_then_table.xlsx": xlsx_title_rows_then_table,
    "xlsx_two_sheets.xlsx": xlsx_two_sheets,
    "xlsx_mixed_units.xlsx": xlsx_mixed_units,
    "xlsx_equivalent_units.xlsx": xlsx_equivalent_units,
    "xlsx_hidden_unit_in_number_format.xlsx": xlsx_missing_units_with_hidden_number_format,
    "xlsx_multiple_test_conditions.xlsx": xlsx_multiple_conditions,
    "xlsx_duplicate_rows.xlsx": xlsx_duplicates,
    "xlsx_conflicting_observations.xlsx": xlsx_conflicting,
    "xlsx_missing_target_column.xlsx": xlsx_missing_target_column,
    "xlsx_data_on_hidden_sheet.xlsx": xlsx_hidden_sheet_only_data,
    "docx_lab_report_two_tables.docx": docx_customer_report,
    "docx_single_table.docx": docx_single_table,
    "docx_conditions_in_table.docx": docx_conditions_in_table,
    "docx_no_tables.docx": docx_no_tables,
    "pdf_supplier_certificate.pdf": pdf_ruled_table,
    "pdf_two_tables.pdf": pdf_two_tables,
    "pdf_table_continued_over_two_pages.pdf": pdf_table_continued_over_two_pages,
    "pdf_scanned_no_text.pdf": pdf_scanned,
}
