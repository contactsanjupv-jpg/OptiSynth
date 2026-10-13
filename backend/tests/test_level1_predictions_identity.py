"""
Level 1 -- the complete predicted-value path, stable candidate identity, no candidate cap,
deterministic ranking and atomic ranking persistence.

Runs through the REAL production path (HTTP routes -> services -> repositories -> SQLite created by
the real Alembic migrations -> engine -> DOCX report). Own isolated database per class.

Run:  python3 -m unittest backend.tests.test_level1_predictions_identity -v
"""
import os
import time
import unittest
from unittest import mock

os.environ.setdefault("SECRET_KEY", "test-secret-key-not-for-production")
os.environ.setdefault("PASSWORD_PEPPER", "test-password-pepper-not-for-production")

import numpy as np  # noqa: E402

from backend.app.services.report_builder import format_predicted_value  # noqa: E402
from backend.tests import evidence_fixtures as F  # noqa: E402
from backend.tests._api_base import SPEC_H, IsolatedApiTestCase  # noqa: E402

HEAD = ["crosslinker_ratio", "cure_temp_c", "salt_spray_hours"]
DECLARED_H = {"declared_units": {"salt_spray_hours": "h"}}  # the file has no unit cells; the reviewer declares hours


def history_csv(n=14):
    return F.csv_bytes(HEAD, F.base_rows(n))


def spread_features(k):
    """Distinct candidate inputs, all inside the historical range (ratio .08-.145, temp 140-179)."""
    return {"crosslinker_ratio": round(0.085 + 0.055 * ((k * 37) % 100) / 99, 5), "cure_temp_c": 141 + (k * 7) % 37}


class Level1Base(IsolatedApiTestCase):
    dir_name = "_test_env_level1"

    def new_case(self, label, n_history=14):
        org_id, client, case_id = self.setup_case(label, SPEC_H)
        self.paid(org_id)
        self.assertEqual(self.upload(client, case_id, "history.csv", history_csv(n_history), DECLARED_H).status_code, 201)
        return org_id, client, case_id

    def add(self, client, case_id, name, features):
        r = client.post(f"/api/change-cases/{case_id}/candidates", json={"name": name, "features": features})
        self.assertEqual(r.status_code, 201, r.text)
        return r.json()

    def rank(self, client, case_id):
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def listing(self, client, case_id):
        r = client.get(f"/api/change-cases/{case_id}/candidates")
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def prediction_rows(self, candidate_ids):
        from sqlalchemy import text
        from backend.app.config import database
        with database.db_connection() as conn:
            return [dict(r) for r in conn.execute(
                text("SELECT * FROM predictions WHERE candidate_id IN :ids ORDER BY id"
                     .replace(":ids", "(" + ",".join(str(int(i)) for i in candidate_ids) + ")"))).mappings().all()]

    def experiment_rows(self, candidate_ids):
        from sqlalchemy import text
        from backend.app.config import database
        with database.db_connection() as conn:
            return [dict(r) for r in conn.execute(
                text("SELECT * FROM recommended_experiments WHERE candidate_id IN :ids ORDER BY id"
                     .replace(":ids", "(" + ",".join(str(int(i)) for i in candidate_ids) + ")"))).mappings().all()]

    def engine_direct(self, candidates, n_history=14):
        """The engine's own answer for the same data, keyed by candidate id."""
        from engine.facade import rank_candidates
        rows = F.base_rows(n_history)
        X = np.array([[r[0], r[1]] for r in rows])
        y = np.array([r[2] for r in rows], dtype=float)
        out = rank_candidates(X, y, "maximize", ["crosslinker_ratio", "cure_temp_c"], 300.0, candidates)
        return {r["id"]: r for r in out}


class TestPredictedValuePath(Level1Base):
    def test_predicted_value_is_persisted_returned_and_matches_the_engine(self):
        org_id, client, case_id = self.new_case("PV1")
        a = self.add(client, case_id, "Low", spread_features(1))
        b = self.add(client, case_id, "High", {"crosslinker_ratio": 0.14, "cure_temp_c": 176})
        ranked = self.rank(client, case_id)
        direct = self.engine_direct([
            {"id": c["id"], "name": c["candidate_name"], "features": c["properties"]}
            for c in self.listing(client, case_id)
        ])
        stored = {p["candidate_id"]: p for p in self.prediction_rows([a["id"], b["id"]])}
        for r in ranked:
            cid = r["candidate_id"]
            self.assertIsNotNone(r["predicted_value"])
            self.assertTrue(np.isfinite(r["predicted_value"]))
            self.assertEqual(stored[cid]["predicted_value"], r["predicted_value"])  # persisted, byte for byte
            self.assertAlmostEqual(r["predicted_value"], direct[cid]["predicted_value"], places=6)
            self.assertAlmostEqual(r["predicted_probability"], direct[cid]["predicted_probability"], places=6)

    def test_value_survives_refresh_and_equals_the_ranking_response(self):
        org_id, client, case_id = self.new_case("PV2")
        for k in range(4):
            self.add(client, case_id, f"C{k}", spread_features(k))
        ranked = self.rank(client, case_id)
        first = self.listing(client, case_id)
        second = self.listing(client, case_id)  # "refresh"
        self.assertEqual(first, second)
        by_id = {c["id"]: c for c in first}
        for r in ranked:
            c = by_id[r["candidate_id"]]
            self.assertEqual(c["candidate_name"], r["candidate_name"])
            self.assertEqual(c["latest_prediction"]["predicted_value"], r["predicted_value"])
            self.assertEqual(c["latest_prediction"]["predicted_probability"], r["predicted_probability"])
            self.assertEqual(c["latest_prediction"]["uncertainty_std"], r["uncertainty_std"])
            self.assertEqual(c["rank"], r["rank"])

    def test_value_is_tied_to_its_dataset_version_and_model_version(self):
        org_id, client, case_id = self.new_case("PV3", n_history=14)
        c = self.add(client, case_id, "Only", spread_features(3))
        self.rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "history2.csv", history_csv(16), DECLARED_H).status_code, 201)
        self.rank(client, case_id)
        d1, d2 = sorted(self.datasets(org_id, case_id), key=lambda d: d["id"])
        rows = self.prediction_rows([c["id"]])
        self.assertEqual([r["dataset_version_id"] for r in rows], [d1["id"], d2["id"]])
        self.assertTrue(all(r["model_version"] == "gp-rank-v1" for r in rows))
        self.assertTrue(all(r["predicted_value"] is not None for r in rows))
        latest = self.listing(client, case_id)[0]["latest_prediction"]
        self.assertEqual(latest["dataset_version_id"], d2["id"])
        self.assertEqual(latest["predicted_value"], rows[-1]["predicted_value"])  # the NEW run's value, not the old one

    def test_prediction_without_a_recorded_value_is_unavailable_never_fabricated(self):
        from backend.app.repositories import candidate_repo, qualification_dataset_repo
        org_id, client, case_id = self.new_case("PV4")
        c = self.add(client, case_id, "Legacy", spread_features(5))
        dataset = qualification_dataset_repo.get_latest_dataset_for_change_case(org_id, case_id)
        # exactly what a prediction stored before this change looks like: no predicted value
        candidate_repo.create_prediction(org_id, c["id"], dataset["id"], "gp-rank-v1", 0.6, 12.0)
        row = self.prediction_rows([c["id"]])[0]
        self.assertIsNone(row["predicted_value"])
        listed = self.listing(client, case_id)[0]
        self.assertIsNone(listed["latest_prediction"]["predicted_value"])
        self.assertEqual(listed["latest_prediction"]["predicted_probability"], 0.6)
        report = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(report.status_code, 200, report.text)
        text = self.docx_text(report.content)
        self.assertIn("Not recorded", text)
        self.assertIn("No value has been estimated or back-filled", text)
        # and still nothing was written behind the user's back
        self.assertIsNone(self.prediction_rows([c["id"]])[0]["predicted_value"])
        # re-running the ranking is the way to obtain a value
        self.rank(client, case_id)
        self.assertIsNotNone(self.listing(client, case_id)[0]["latest_prediction"]["predicted_value"])


class TestReportCarriesTheSameResults(Level1Base):
    def test_docx_shows_value_requirement_direction_unit_uncertainty_probability_and_warnings(self):
        org_id, client, case_id = self.new_case("RP1")
        self.add(client, case_id, "Inside", spread_features(2))
        self.add(client, case_id, "Far outside", {"crosslinker_ratio": 0.5, "cure_temp_c": 400})
        ranked = self.rank(client, case_id)
        report = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(report.status_code, 200, report.text)
        text = self.docx_text(report.content)
        self.assertIn("Requirement: salt_spray_hours (h) at or above 300.0 (direction: maximize)", text)
        self.assertIn("MODEL ESTIMATES, not measured physical results", text)
        self.assertIn("Predicted salt_spray_hours (h) (model estimate)", text)
        self.assertIn("Uncertainty (std, h)", text)
        self.assertIn("Model-estimated probability", text)
        for r in ranked:
            self.assertIn(format_predicted_value(r["predicted_value"]), text)
            self.assertIn(f"{r['predicted_probability']:.0%}", text)
            self.assertIn(f"{r['uncertainty_std']:.3f}", text)
        self.assertIn("Outside historical range", text)      # domain warning
        self.assertIn("Domain caution", text)                 # recommended-experiment warning text
        self.assertIn("Evidence intake review", text)         # data provenance
        self.assertIn("history.csv", text)
        self.assertIn("Recommended Validation Experiments", text)
        self.assertIn("equal probability are ordered by predicted value", text)

    def test_docx_rank_order_matches_the_ranking_response_and_the_listing(self):
        from docx import Document
        import io
        org_id, client, case_id = self.new_case("RP2")
        for k in range(7):
            self.add(client, case_id, f"Cand {k}", spread_features(k))
        ranked = self.rank(client, case_id)
        report = client.post(f"/api/change-cases/{case_id}/report")
        doc = Document(io.BytesIO(report.content))
        table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Rank")
        names = [row.cells[1].text for row in table.rows[1:]]
        self.assertEqual(names, [r["candidate_name"] for r in ranked])
        self.assertEqual([row.cells[0].text for row in table.rows[1:]], [str(i) for i in range(1, 8)])
        listed = sorted(self.listing(client, case_id), key=lambda c: c["rank"])
        self.assertEqual([c["candidate_name"] for c in listed], names)

    def test_report_without_a_declared_target_unit_says_so(self):
        from backend.tests._api_base import SPEC
        org_id, client, case_id = self.setup_case("RP3", SPEC)
        self.paid(org_id)
        self.assertEqual(self.upload(client, case_id, "history.csv", history_csv()).status_code, 201)
        self.add(client, case_id, "A", spread_features(1))
        self.rank(client, case_id)
        text = self.docx_text(client.post(f"/api/change-cases/{case_id}/report").content)
        self.assertIn("No canonical unit is declared for this property", text)
        self.assertIn("Predicted salt_spray_hours (model estimate)", text)


class TestCandidateIdentity(Level1Base):
    def _twins(self, label):
        """Two candidates with the SAME name but different inputs, plus one other. The API refuses
        duplicate names, so the twins are inserted straight through the repository -- the situation
        the ranking must survive even if a duplicate ever reaches the database."""
        import json
        from backend.app.repositories import candidate_repo
        org_id, client, case_id = self.new_case(label)
        low = candidate_repo.create_candidate(org_id, case_id, "Twin", {"crosslinker_ratio": 0.09, "cure_temp_c": 143})
        high = candidate_repo.create_candidate(org_id, case_id, "Twin", {"crosslinker_ratio": 0.14, "cure_temp_c": 176})
        solo = self.add(client, case_id, "Solo", {"crosslinker_ratio": 0.11, "cure_temp_c": 160})["id"]
        return org_id, client, case_id, low, high, solo

    def test_duplicate_names_cannot_attach_a_prediction_to_the_wrong_candidate(self):
        org_id, client, case_id, low, high, solo = self._twins("ID1")
        ranked = self.rank(client, case_id)
        self.assertEqual(sorted(r["candidate_id"] for r in ranked), sorted([low, high, solo]))
        rows = self.prediction_rows([low, high, solo])
        self.assertEqual(sorted(r["candidate_id"] for r in rows), sorted([low, high, solo]))  # exactly one each
        self.assertEqual(len(self.experiment_rows([low, high, solo])), 3)
        direct = self.engine_direct([
            {"id": c["id"], "name": c["candidate_name"], "features": c["properties"]}
            for c in self.listing(client, case_id)
        ])
        stored = {r["candidate_id"]: r for r in rows}
        for cid in (low, high, solo):
            self.assertAlmostEqual(stored[cid]["predicted_value"], direct[cid]["predicted_value"], places=6, msg=cid)
            self.assertAlmostEqual(stored[cid]["predicted_probability"], direct[cid]["predicted_probability"], places=6)
        # different inputs really did produce different estimates for the two same-named candidates
        self.assertNotAlmostEqual(stored[low]["predicted_value"], stored[high]["predicted_value"], places=3)

    def test_listing_attaches_each_prediction_to_its_own_candidate(self):
        org_id, client, case_id, low, high, solo = self._twins("ID2")
        ranked = {r["candidate_id"]: r for r in self.rank(client, case_id)}
        for c in self.listing(client, case_id):
            self.assertEqual(c["latest_prediction"]["predicted_value"], ranked[c["id"]]["predicted_value"])

    def test_report_keeps_same_named_candidates_apart(self):
        from docx import Document
        import io
        org_id, client, case_id, low, high, solo = self._twins("ID3")
        ranked = self.rank(client, case_id)
        doc = Document(io.BytesIO(client.post(f"/api/change-cases/{case_id}/report").content))
        table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Rank")
        shown = [(row.cells[1].text, row.cells[2].text) for row in table.rows[1:]]
        self.assertEqual(shown, [(r["candidate_name"], format_predicted_value(r["predicted_value"])) for r in ranked])

    def test_identical_inputs_tie_and_order_by_id_whatever_the_names(self):
        org_id, client, case_id = self.new_case("ID4")
        feats = {"crosslinker_ratio": 0.11, "cure_temp_c": 160}
        first = self.add(client, case_id, "Zulu", feats)["id"]
        second = self.add(client, case_id, "Alpha", feats)["id"]
        third = self.add(client, case_id, "Mike", feats)["id"]
        a = [r["candidate_id"] for r in self.rank(client, case_id)]
        b = [r["candidate_id"] for r in self.rank(client, case_id)]
        self.assertEqual(a, [first, second, third])
        self.assertEqual(a, b)
        self.assertEqual([c["rank"] for c in self.listing(client, case_id)], [1, 2, 3])

    def test_api_rejects_duplicate_names_trimmed_and_case_insensitive(self):
        org_id, client, case_id = self.new_case("ID5")
        self.add(client, case_id, "Resin A", spread_features(1))
        for variant in ("Resin A", "RESIN a", "   resin a "):
            r = client.post(f"/api/change-cases/{case_id}/candidates",
                            json={"name": variant, "features": spread_features(2)})
            self.assertEqual(r.status_code, 400, variant)
            self.assertIn("already exists", r.json()["error"])
        self.assertEqual(len(self.listing(client, case_id)), 1)

    def test_same_name_is_allowed_in_a_different_diagnostic(self):
        org_id, client, case_id = self.new_case("ID6")
        other = self.make_case(client, SPEC_H)
        self.add(client, case_id, "Resin A", spread_features(1))
        self.add(client, other, "Resin A", spread_features(1))  # different diagnostic: fine


class TestCandidateCountScale(Level1Base):
    """No product cap. Counts 1, 5, 6, 25 and 100; ranking order, identities, prediction
    associations and timing are checked at each size."""
    TIMINGS = {}

    def _run(self, n):
        from backend.app.repositories import candidate_repo
        from engine.facade import ranking_sort_key
        org_id, client, case_id = self.new_case(f"Scale{n}")
        t0 = time.perf_counter()
        ids = [self.add(client, case_id, f"Candidate {k:03d}", spread_features(k))["id"] for k in range(n)]
        add_seconds = time.perf_counter() - t0
        self.assertEqual(len(set(ids)), n)

        persist_seconds = []
        real = candidate_repo.persist_ranking

        def timed(*a, **kw):
            t = time.perf_counter()
            try:
                return real(*a, **kw)
            finally:
                persist_seconds.append(time.perf_counter() - t)

        t0 = time.perf_counter()
        with mock.patch.object(candidate_repo, "persist_ranking", side_effect=timed):
            ranked = self.rank(client, case_id)
        rank_seconds = time.perf_counter() - t0

        self.assertEqual(len(ranked), n)
        self.assertEqual(sorted(r["candidate_id"] for r in ranked), sorted(ids))
        self.assertEqual([r["rank"] for r in ranked], list(range(1, n + 1)))
        keys = [ranking_sort_key(r["predicted_probability"], r["predicted_value"], r["candidate_id"], "maximize")
                for r in ranked]
        self.assertEqual(keys, sorted(keys))
        # persisted: exactly one prediction and one recommended experiment per candidate, correct association
        rows = self.prediction_rows(ids)
        self.assertEqual(sorted(r["candidate_id"] for r in rows), sorted(ids))
        self.assertEqual(sorted(r["candidate_id"] for r in self.experiment_rows(ids)), sorted(ids))
        by_id = {r["candidate_id"]: r for r in rows}
        for r in ranked:
            self.assertEqual(by_id[r["candidate_id"]]["predicted_value"], r["predicted_value"])
        # retrieval agrees with the ranking run
        listed = self.listing(client, case_id)
        self.assertEqual([c["id"] for c in listed], sorted(ids))  # listing stays in id order
        self.assertEqual({c["id"]: c["rank"] for c in listed}, {r["candidate_id"]: r["rank"] for r in ranked})
        self.assertEqual([c["id"] for c in sorted(listed, key=lambda c: c["rank"])],
                         [r["candidate_id"] for r in ranked])
        type(self).TIMINGS[n] = (add_seconds, rank_seconds, persist_seconds[0])
        print(f"\n[scale] n={n:>3}: add {add_seconds:6.2f}s total | rank (incl. persist) {rank_seconds:6.2f}s | "
              f"persist {persist_seconds[0] * 1000:7.1f} ms", end="")
        return ranked

    def test_1_candidate(self):
        self._run(1)

    def test_5_candidates(self):
        self._run(5)

    def test_6_candidates_one_past_the_old_cap(self):
        self._run(6)

    def test_25_candidates(self):
        self._run(25)

    def test_100_candidates(self):
        self._run(100)
        # A generous guard against a pathological regression only -- NOT a product limit.
        add_s, rank_s, persist_s = type(self).TIMINGS[100]
        self.assertLess(rank_s, 120)
        self.assertLess(persist_s, 30)

    def test_report_generates_for_100_candidates_in_rank_order(self):
        from docx import Document
        import io
        org_id, client, case_id = self.new_case("ScaleReport")
        for k in range(100):
            self.add(client, case_id, f"Candidate {k:03d}", spread_features(k))
        ranked = self.rank(client, case_id)
        resp = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(resp.status_code, 200, resp.text[:200])
        doc = Document(io.BytesIO(resp.content))
        table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Rank")
        self.assertEqual([row.cells[1].text for row in table.rows[1:]], [r["candidate_name"] for r in ranked])


class TestAtomicRankingPersistence(Level1Base):
    def _counts(self, ids):
        return len(self.prediction_rows(ids)), len(self.experiment_rows(ids))

    def test_repository_rolls_back_the_whole_run_when_one_row_is_bad(self):
        from sqlalchemy.exc import IntegrityError
        from backend.app.repositories import candidate_repo, qualification_dataset_repo
        org_id, client, case_id = self.new_case("AT1")
        ids = [self.add(client, case_id, f"C{k}", spread_features(k))["id"] for k in range(4)]
        dataset = qualification_dataset_repo.get_latest_dataset_for_change_case(org_id, case_id)

        def row(cid, prob):
            return {"candidate_id": cid, "model_version": "gp-rank-v1", "predicted_value": 500.0,
                    "predicted_probability": prob, "uncertainty_std": 10.0,
                    "experiment_description": "run it", "priority": 1}

        rows = [row(ids[0], 0.9), row(ids[1], 0.8), row(ids[2], None), row(ids[3], 0.6)]  # NOT NULL violation on #3
        with self.assertRaises(IntegrityError):
            candidate_repo.persist_ranking(org_id, dataset["id"], rows)
        self.assertEqual(self._counts(ids), (0, 0))  # not even the two valid rows before it were kept
        candidate_repo.persist_ranking(org_id, dataset["id"], [row(i, 0.5) for i in ids])
        self.assertEqual(self._counts(ids), (4, 4))

    def test_service_failure_midway_leaves_the_previous_ranking_untouched(self):
        from backend.app.repositories import candidate_repo
        from backend.app.services import change_case_service
        org_id, client, case_id = self.new_case("AT2")
        ids = [self.add(client, case_id, f"C{k}", spread_features(k))["id"] for k in range(5)]
        self.rank(client, case_id)
        before = self.listing(client, case_id)
        self.assertEqual(self._counts(ids), (5, 5))

        real = candidate_repo._insert_recommended_experiment
        calls = {"n": 0}

        def flaky(*a, **kw):
            calls["n"] += 1
            if calls["n"] == 3:
                raise RuntimeError("simulated database failure")
            return real(*a, **kw)

        with mock.patch.object(candidate_repo, "_insert_recommended_experiment", side_effect=flaky):
            with self.assertRaises(RuntimeError):
                change_case_service.rank_change_case(org_id, case_id)
        self.assertEqual(calls["n"], 3)
        self.assertEqual(self._counts(ids), (5, 5))           # nothing from the failed run survived
        self.assertEqual(self.listing(client, case_id), before)  # customers still see the earlier ranking intact

    def test_prediction_insert_failure_is_also_all_or_nothing(self):
        from backend.app.repositories import candidate_repo
        from backend.app.services import change_case_service
        org_id, client, case_id = self.new_case("AT3")
        ids = [self.add(client, case_id, f"C{k}", spread_features(k))["id"] for k in range(4)]
        real = candidate_repo._insert_prediction
        calls = {"n": 0}

        def flaky(*a, **kw):
            calls["n"] += 1
            if calls["n"] == 2:
                raise RuntimeError("simulated database failure")
            return real(*a, **kw)

        with mock.patch.object(candidate_repo, "_insert_prediction", side_effect=flaky):
            with self.assertRaises(RuntimeError):
                change_case_service.rank_change_case(org_id, case_id)
        self.assertEqual(self._counts(ids), (0, 0))
        self.assertTrue(all(c["latest_prediction"] is None and c["rank"] is None for c in self.listing(client, case_id)))
        # and a clean retry works
        self.rank(client, case_id)
        self.assertEqual(self._counts(ids), (4, 4))

    def test_engine_returning_the_wrong_candidates_saves_nothing(self):
        from backend.app.services import change_case_service
        org_id, client, case_id = self.new_case("AT4")
        ids = [self.add(client, case_id, f"C{k}", spread_features(k))["id"] for k in range(3)]

        def broken(X, y, direction, cols, target, candidates, **kw):
            return [{**c, "id": c["id"] + 1000, "predicted_value": 1.0, "uncertainty_std": 1.0,
                     "predicted_probability": 0.5, "model_version": "x"} for c in candidates]

        with mock.patch.object(change_case_service, "engine_rank_candidates", side_effect=broken):
            with self.assertRaises(RuntimeError):
                change_case_service.rank_change_case(org_id, case_id)
        self.assertEqual(self._counts(ids), (0, 0))


class TestPredictedValueMigration(Level1Base):
    dir_name = "_test_env_level1_migration"

    def _columns(self):
        from sqlalchemy import text
        from backend.app.config import database
        with database.db_connection() as conn:
            return [r[1] for r in conn.execute(text("PRAGMA table_info(predictions)")).all()]

    def _alembic(self, direction, revision):
        from alembic.config import Config
        from alembic import command
        backend_dir = os.path.join(os.path.dirname(__file__), "..")
        cfg = Config(os.path.join(backend_dir, "alembic.ini"))
        cfg.set_main_option("script_location", os.path.join(backend_dir, "migrations"))
        getattr(command, direction)(cfg, revision)

    def test_migration_is_additive_nullable_and_reversible_without_touching_other_data(self):
        from backend.app.repositories import candidate_repo, qualification_dataset_repo
        org_id, client, case_id = self.new_case("MG1")
        c = self.add(client, case_id, "Legacy", spread_features(4))
        dataset = qualification_dataset_repo.get_latest_dataset_for_change_case(org_id, case_id)
        candidate_repo.create_prediction(org_id, c["id"], dataset["id"], "gp-rank-v1", 0.42, 7.5)  # no value
        self.assertIn("predicted_value", self._columns())

        self._alembic("downgrade", "b8d52f1a3c77")
        self.assertNotIn("predicted_value", self._columns())
        survivors = self.prediction_rows([c["id"]])
        self.assertEqual((survivors[0]["predicted_probability"], survivors[0]["uncertainty_std"]), (0.42, 7.5))

        self._alembic("upgrade", "head")
        self.assertIn("predicted_value", self._columns())
        rows = self.prediction_rows([c["id"]])
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0]["predicted_value"])  # never back-filled
        self.assertEqual((rows[0]["predicted_probability"], rows[0]["uncertainty_std"]), (0.42, 7.5))

        # and the application works against the re-upgraded schema
        self.rank(client, case_id)
        self.assertIsNotNone(self.listing(client, case_id)[0]["latest_prediction"]["predicted_value"])


if __name__ == "__main__":
    unittest.main()
