"""
Ranking requires an explicit, valid canonical unit for EVERY modelled input feature and for the target.

THE DEFECT. The sufficiency check "UNITS_DECLARED" only warned (and only for datasets that had an intake review
record), and the ranking run never consulted it at all: a diagnostic whose fields had no declared unit could be
ranked, so candidate values and evidence were never checked as comparable.

THE BEHAVIOUR. rank_change_case() refuses -- before anything is computed or saved -- naming each field that needs a
unit and how to resolve it. Units are never guessed from names or values; a dimensionless ratio, a fraction and a
percentage are declared explicitly ('dimensionless', 'fraction', '%'), and fraction and % are never converted into
each other. A diagnostic's units cannot be edited, so the resolution for an existing diagnostic is a new one.
Legacy diagnostics (no units, possibly already ranked) are left exactly as stored, flagged in the evidence check and
in the report, and refused for any NEW ranking.

Synthetic data: these tests prove the gate and its messages, not the science.

Run:  python3 -m unittest backend.tests.test_units_before_ranking -v
"""
import io
import json
import unittest

from backend.tests import evidence_fixtures as F
from backend.tests._api_base import SPEC, SPEC_NO_UNITS, IsolatedApiTestCase, persist_legacy_ranking

IN_A = {"crosslinker_ratio": 0.12, "cure_temp_c": 158}
IN_B = {"crosslinker_ratio": 0.10, "cure_temp_c": 150}
FAR = {"crosslinker_ratio": 0.40, "cure_temp_c": 300}


class TestUnitsRule(unittest.TestCase):
    def test_every_field_and_the_target_are_checked_and_named(self):
        from backend.app.services import change_case_rules as R
        problems = R.units_declaration_problems(SPEC_NO_UNITS)
        self.assertEqual([(p["field"], p["role"]) for p in problems],
                         [("crosslinker_ratio", "input feature"), ("cure_temp_c", "input feature"),
                          ("salt_spray_hours", "target")])
        self.assertEqual(R.units_declaration_problems(SPEC), [])

    def test_explicit_dimensionless_representations_are_valid_and_blank_is_not(self):
        from backend.app.services import change_case_rules as R
        for token in ("dimensionless", "fraction", "%", "wt%", "1", "score"):  # three explicit + custom labels
            spec = {**SPEC, "units": {**SPEC["units"], "crosslinker_ratio": token}}
            self.assertEqual(R.units_declaration_problems(spec), [], token)
        for blank in ("", "   ", None):
            spec = {**SPEC, "units": {**SPEC["units"], "crosslinker_ratio": blank}}
            self.assertEqual([p["field"] for p in R.units_declaration_problems(spec)], ["crosslinker_ratio"], repr(blank))

    def test_a_unit_that_differs_from_a_supported_one_only_by_case_is_ambiguous(self):
        from backend.app.services import change_case_rules as R
        spec = {**SPEC, "units": {**SPEC["units"], "salt_spray_hours": "H"}}
        problems = R.units_declaration_problems(spec)
        self.assertEqual(len(problems), 1)
        self.assertIn("ambiguous", problems[0]["problem"])
        self.assertIn("'h'", problems[0]["problem"])

    def test_the_check_is_blocking_with_and_without_an_intake_review_record(self):
        from backend.app.services import change_case_rules as R
        rows = [{"features": {"crosslinker_ratio": a, "cure_temp_c": b}, "target_value": c} for a, b, c in F.base_rows(14)]
        cand = [{"candidate_name": "A", "properties": IN_A}]
        for review in ({"flags": []}, None):  # None = a legacy dataset with no intake review record
            result = R.assess_evidence_sufficiency(SPEC_NO_UNITS, rows, cand, True, review)
            self.assertIn("UNITS_DECLARED", [c["code"] for c in result["blocking"]])
            self.assertFalse(result["can_rank"])

    def test_the_message_names_fields_and_explains_the_resolution(self):
        from backend.app.schemas.errors import ValidationError
        from backend.app.services import change_case_rules as R
        with self.assertRaises(ValidationError) as ctx:
            R.check_units_declared_for_ranking(SPEC_NO_UNITS)
        message = ctx.exception.message
        for needle in ("crosslinker_ratio (input feature)", "cure_temp_c (input feature)", "salt_spray_hours (target)",
                       "never guessed", "'fraction'", "'%'", "'dimensionless'", "never converted into each other",
                       "create a new diagnostic"):
            self.assertIn(needle, message)


class TestRankingRequiresUnitsApi(IsolatedApiTestCase):
    dir_name = "_test_env_units_rank"

    def _case(self, label, spec, upload=True, candidates=(("InA", IN_A),)):
        org_id, client, case_id = self.setup_case(label, spec)
        self.paid(org_id)
        if upload:
            up = self.upload(client, case_id, "history.csv", F.csv_bytes(F.CANON_HEADERS, F.base_rows(14)))
            self.assertEqual(up.status_code, 201, up.text)
        ids = {}
        for name, features in candidates:
            r = client.post(f"/api/change-cases/{case_id}/candidates", json={"name": name, "features": features})
            self.assertEqual(r.status_code, 201, r.text)
            ids[name] = r.json()["id"]
        return org_id, client, case_id, ids

    def _listing(self, client, case_id):
        return {c["candidate_name"]: c for c in client.get(f"/api/change-cases/{case_id}/candidates").json()}

    def _report(self, client, case_id):
        from docx import Document
        r = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(r.status_code, 200, r.text[:200])
        doc = Document(io.BytesIO(r.content))
        return [p.text for p in doc.paragraphs], [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]

    # ---------------------------------------------------------------- blocked
    def test_ranking_without_any_declared_unit_is_refused_naming_every_field(self):
        org_id, client, case_id, ids = self._case("blocked", SPEC_NO_UNITS)
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400, r.text)
        message = r.json()["error"]
        for field in ("crosslinker_ratio", "cure_temp_c", "salt_spray_hours"):
            self.assertIn(f"{field} (", message)
        self.assertIn("never guessed", message)
        self.assertIn("create a new diagnostic", message)

    def test_a_refused_ranking_persists_nothing_and_the_report_does_not_imply_it_ran(self):
        from backend.app.repositories import candidate_repo
        org_id, client, case_id, ids = self._case("nothing_saved", SPEC_NO_UNITS)
        before = candidate_repo.get_candidate_usage_map(org_id, case_id)
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 400)
        after = candidate_repo.get_candidate_usage_map(org_id, case_id)
        self.assertEqual(before, after)
        self.assertEqual(after[ids["InA"]], {"prediction_count": 0, "outcome_count": 0, "experiment_count": 0})
        self.assertIsNone(self._listing(client, case_id)["InA"]["latest_prediction"])
        paragraphs, tables = self._report(client, case_id)
        self.assertIn("No candidates ranked yet.", paragraphs)
        self.assertFalse([t for t in tables if t[0][0] == "Rank"], "no ranked table may appear")

    def test_the_units_are_checked_before_a_dataset_is_even_looked_at(self):
        org_id, client, case_id, _ = self._case("no_dataset", SPEC_NO_UNITS, upload=False)
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400)
        self.assertIn("explicit canonical unit", r.json()["error"])

    def test_the_evidence_check_reports_the_block_before_anyone_tries_to_rank(self):
        _, client, case_id, _ = self._case("sufficiency", SPEC_NO_UNITS)
        body = client.get(f"/api/change-cases/{case_id}/sufficiency").json()
        units = next(c for c in body["checks"] if c["code"] == "UNITS_DECLARED")
        self.assertEqual(units["status"], "fail")
        self.assertFalse(body["can_rank"])

    def test_only_the_fields_without_a_valid_unit_are_named(self):
        partial = {**SPEC_NO_UNITS, "units": {"cure_temp_c": "degC", "salt_spray_hours": "h"}}
        _, client, case_id, _ = self._case("partial", partial)
        message = client.post(f"/api/change-cases/{case_id}/rank").json()["error"]
        self.assertIn("crosslinker_ratio (input feature): no unit declared", message)
        self.assertNotIn("cure_temp_c (", message)
        self.assertNotIn("salt_spray_hours (", message)

    def test_an_ambiguous_unit_is_refused_not_silently_accepted(self):
        spec = {**SPEC, "units": {**SPEC["units"], "salt_spray_hours": "H"}}
        _, client, case_id, _ = self._case("ambiguous", spec)
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400)
        self.assertIn("'H' is ambiguous", r.json()["error"])

    # ---------------------------------------------------------------- allowed
    def test_ranking_succeeds_once_every_unit_is_declared_and_saves_the_prediction(self):
        from backend.app.repositories import candidate_repo
        org_id, client, case_id, ids = self._case("ok", SPEC, candidates=(("InA", IN_A), ("InB", IN_B)))
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(sorted(x["rank"] for x in r.json()), [1, 2])
        usage = candidate_repo.get_candidate_usage_map(org_id, case_id)
        self.assertTrue(all(u["prediction_count"] == 1 for u in usage.values()))
        sufficiency = client.get(f"/api/change-cases/{case_id}/sufficiency").json()
        self.assertEqual(next(c for c in sufficiency["checks"] if c["code"] == "UNITS_DECLARED")["status"], "pass")

    def test_fraction_and_percent_are_both_valid_declarations_but_are_never_converted_into_each_other(self):
        spec = {**SPEC, "units": {**SPEC["units"], "crosslinker_ratio": "%"}}
        _, client, case_id, _ = self._case("percent", spec, candidates=())
        url = f"/api/change-cases/{case_id}/candidates"
        wrong = client.post(url, json={"name": "X", "features": IN_A, "units": {"crosslinker_ratio": "fraction", "cure_temp_c": "degC"}})
        self.assertEqual(wrong.status_code, 400)  # a fraction is not silently turned into a percentage
        right = client.post(url, json={"name": "X", "features": IN_A, "units": {"crosslinker_ratio": "%", "cure_temp_c": "degC"}})
        self.assertEqual(right.status_code, 201, right.text)
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 200)

    # ---------------------------------------------------------------- legacy diagnostics
    def test_a_legacy_diagnostic_is_not_rewritten_not_guessed_and_not_rankable_but_stays_readable(self):
        from backend.app.config.database import db_connection
        from backend.app.repositories import candidate_repo, change_case_repo
        from sqlalchemy import text
        org_id, client, case_id, ids = self._case("legacy", SPEC_NO_UNITS, candidates=(("InA", IN_A), ("InB", IN_B)))
        persist_legacy_ranking(org_id, case_id)   # a ranking stored before units were required

        def snapshot():
            with db_connection() as conn:
                experiments = conn.execute(text(
                    "SELECT features_json, target_value FROM qualification_experiments WHERE change_case_id = :c ORDER BY id"),
                    {"c": case_id}).all()
            return (change_case_repo.get_change_case(org_id, case_id)["qualification_spec_json"],
                    [tuple(r) for r in experiments], candidate_repo.get_candidate_usage_map(org_id, case_id),
                    [(c["id"], c["properties_json"], c.get("input_json")) for c in
                     candidate_repo.list_candidates_for_change_case(org_id, case_id)])

        before = snapshot()
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 400)
        self.assertIn("create a new diagnostic", r.json()["error"])
        self.assertEqual(snapshot(), before, "stored spec, evidence rows, predictions and candidates are untouched")
        self.assertNotIn('"units"', before[0], "no unit was invented for the legacy spec")
        # the earlier results stay readable
        self.assertIsNotNone(self._listing(client, case_id)["InA"]["latest_prediction"])
        paragraphs, _ = self._report(client, case_id)
        text_ = " ".join(paragraphs)
        self.assertIn("Canonical units were not declared for: crosslinker_ratio, cure_temp_c, salt_spray_hours", text_)
        self.assertIn("produced without unit checking", text_)

    # ---------------------------------------------------------------- earlier protections are intact
    def test_the_stale_report_refusal_still_applies_to_a_unit_complete_ranking(self):
        _, client, case_id, _ = self._case("fix1", SPEC, candidates=(("InA", IN_A), ("InB", IN_B)))
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/rank").status_code, 200)
        self.assertEqual(client.post(f"/api/change-cases/{case_id}/report").status_code, 200)
        client.post(f"/api/change-cases/{case_id}/candidates", json={"name": "InC", "features": {"crosslinker_ratio": 0.13, "cure_temp_c": 165}})
        refused = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(refused.status_code, 400)
        self.assertIn("out of date", refused.json()["error"])

    def test_the_domain_aware_safeguard_still_leaves_an_outside_candidate_unranked(self):
        _, client, case_id, _ = self._case("domain", SPEC, candidates=(("InA", IN_A), ("InB", IN_B), ("Far", FAR)))
        by = {x["candidate_name"]: x for x in client.post(f"/api/change-cases/{case_id}/rank").json()}
        self.assertIsNone(by["Far"]["rank"])
        self.assertEqual(by["Far"]["ranking_status"], "not_ranked_insufficient_evidence")
        self.assertEqual(sorted(x["rank"] for n, x in by.items() if n != "Far"), [1, 2])
        for key in ("r2_score", "accuracy_tolerance_note"):   # the metric changes are intact
            self.assertIn(key, by["InA"]["model_quality"])


if __name__ == "__main__":
    unittest.main()
