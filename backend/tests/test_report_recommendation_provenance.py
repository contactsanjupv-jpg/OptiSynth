"""
Report provenance regression: the report's "Recommended Validation Experiments" must be
built from the SAME latest prediction that the "Ranked Candidates" table shows.

The bug this guards against: every ranking run stores a new recommended_experiments row
(all with priority 1) and the report used the first row returned by
`ORDER BY priority DESC, id ASC` -- the OLDEST. After a second ranking on different
evidence the Ranked Candidates table showed the current probabilities while the
Recommended Validation Experiments still quoted the first ranking's.

Two (or three) ranking runs with deliberately DIFFERENT results are performed, and the
numbers in the generated DOCX are compared exactly.

Run:  python3 -m unittest backend.tests.test_report_recommendation_provenance -v
"""
import io
import re
import unittest

from backend.tests import evidence_fixtures as F
from backend.tests._api_base import IsolatedApiTestCase


def _dataset(target_of_i, n=14):
    """Same 14 input rows every time; only the TARGET values differ between datasets."""
    return F.csv_bytes(F.CANON_HEADERS, [(a, b, target_of_i(i)) for i, (a, b, _) in enumerate(F.base_rows(n))])


HIGH = _dataset(lambda i: 100 + 40 * i)    # targets 100..620  -> mid-range candidates meet the 300 target (~100%)
LOW = _dataset(lambda i: 100 + 10 * i)     # targets 100..230  -> nothing meets 300 (~0%)
MIDDLE = _dataset(lambda i: 100 + 20 * i)  # targets 100..360  -> in between: a third, different set of results


class TestReportRecommendationsUseTheLatestRanking(IsolatedApiTestCase):
    dir_name = "_test_env_recprov"

    CANDIDATES = (("Mid", 0.11, 155), ("Upper", 0.13, 168), ("Beyond", 0.16, 178))  # Beyond is outside both ranges

    def _setup(self, label):
        org_id, client, case_id = self.setup_case(label)
        self.paid(org_id)
        self.assertEqual(self.upload(client, case_id, "first.csv", HIGH).status_code, 201)
        for name, ratio, temp in self.CANDIDATES:
            r = client.post(f"/api/change-cases/{case_id}/candidates",
                            json={"name": name, "features": {"crosslinker_ratio": ratio, "cure_temp_c": temp}})
            self.assertEqual(r.status_code, 201, r.text)
        return org_id, client, case_id

    def _rank(self, client, case_id):
        r = client.post(f"/api/change-cases/{case_id}/rank")
        self.assertEqual(r.status_code, 200, r.text)
        return {x["candidate_name"]: x for x in r.json()}

    @staticmethod
    def _pct(p):
        return f"{p:.0%}"

    def _report(self, client, case_id):
        """-> (ranked_table {name: 'NN%'}, recommendations {name: full text}) read from the real DOCX."""
        from docx import Document
        resp = client.post(f"/api/change-cases/{case_id}/report")
        self.assertEqual(resp.status_code, 200, resp.text)
        doc = Document(io.BytesIO(resp.content))
        table = next(t for t in doc.tables if t.rows[0].cells[1].text == "Model-estimated probability")
        ranked = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
        recs, on = {}, False
        for p in doc.paragraphs:
            if p.style.name.startswith("Heading"):
                on = p.text == "Recommended Validation Experiments"
            elif on and p.text.strip():
                name, _, text = p.text.partition(": ")
                recs[name] = text
        return ranked, recs

    def _estimate_in(self, text):
        m = re.search(r"model estimate (\d+)%", text)
        self.assertIsNotNone(m, text)
        return f"{m.group(1)}%"

    # ------------------------------------------------------------------
    def test_the_two_rankings_really_differ_so_the_test_is_meaningful(self):
        org_id, client, case_id = self._setup("RecPremise")
        first = self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        second = self._rank(client, case_id)
        for name in ("Mid", "Upper"):
            self.assertNotEqual(self._pct(first[name]["predicted_probability"]),
                                self._pct(second[name]["predicted_probability"]), name)

    def test_after_two_rankings_the_report_table_and_recommendations_show_the_same_latest_numbers(self):
        org_id, client, case_id = self._setup("RecTwo")
        first = self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        second = self._rank(client, case_id)

        ranked, recs = self._report(client, case_id)
        self.assertEqual(set(ranked), {"Mid", "Upper", "Beyond"})
        for name in ranked:
            latest = self._pct(second[name]["predicted_probability"])
            older = self._pct(first[name]["predicted_probability"])
            self.assertEqual(ranked[name], latest, f"{name}: ranked table must show the latest ranking")
            self.assertEqual(self._estimate_in(recs[name]), latest,
                             f"{name}: recommendation must quote the SAME latest probability")
            if older != latest:
                self.assertNotIn(f"model estimate {older}", recs[name], f"{name}: must not quote the first ranking")

    def test_the_old_recommendations_are_still_stored_only_the_report_no_longer_reads_them(self):
        """Proves the failing condition existed (two stored recommendations per candidate, the oldest with
        the first ranking's number) -- and that storage is untouched by the fix."""
        org_id, client, case_id = self._setup("RecHistory")
        first = self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        second = self._rank(client, case_id)
        from backend.app.config.database import db_connection
        from sqlalchemy import text
        with db_connection() as conn:
            rows = conn.execute(text(
                "SELECT c.candidate_name, r.id, r.description FROM recommended_experiments r "
                "JOIN candidate_substitutes c ON c.id = r.candidate_id WHERE c.change_case_id = :cid ORDER BY r.id"),
                {"cid": case_id}).all()
        by = {}
        for name, rid, desc in rows:
            by.setdefault(name, []).append(desc)
        for name in ("Mid", "Upper"):
            self.assertEqual(len(by[name]), 2, name)  # history kept
            self.assertEqual(self._estimate_in(by[name][0]), self._pct(first[name]["predicted_probability"]))
            self.assertEqual(self._estimate_in(by[name][1]), self._pct(second[name]["predicted_probability"]))

    def test_three_rankings_the_report_follows_the_third(self):
        org_id, client, case_id = self._setup("RecThree")
        first = self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        second = self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "third.csv", MIDDLE).status_code, 201)
        third = self._rank(client, case_id)
        # premise: for "Upper" all three runs gave different numbers, so neither the oldest nor the
        # middle stored recommendation can be mistaken for the latest
        pcts = [self._pct(r["Upper"]["predicted_probability"]) for r in (first, second, third)]
        self.assertEqual(len(set(pcts)), 3, pcts)
        ranked, recs = self._report(client, case_id)
        for name in ranked:
            latest = self._pct(third[name]["predicted_probability"])
            self.assertEqual((ranked[name], self._estimate_in(recs[name])), (latest, latest), name)

    def test_before_re_ranking_the_report_is_consistent_with_the_ranking_it_actually_describes(self):
        """New evidence accepted but NOT re-ranked: the report still describes ranking #1 -- and BOTH
        sections must say so, not one section from ranking #1 and another from anything else."""
        org_id, client, case_id = self._setup("RecStale")
        first = self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        ranked, recs = self._report(client, case_id)
        for name in ranked:
            expected = self._pct(first[name]["predicted_probability"])
            self.assertEqual((ranked[name], self._estimate_in(recs[name])), (expected, expected), name)

    def test_domain_caution_in_the_recommendation_follows_the_latest_coverage(self):
        org_id, client, case_id = self._setup("RecCoverage")
        self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        self._rank(client, case_id)
        cands = {c["candidate_name"]: c for c in client.get(f"/api/change-cases/{case_id}/candidates").json()}
        _, recs = self._report(client, case_id)
        for name, c in cands.items():
            expect_caution = c["domain_coverage"]["status"] != "within_historical_domain"
            self.assertEqual("Domain caution" in recs[name], expect_caution, name)
        self.assertIn("Domain caution", recs["Beyond"])

    def test_tier_wording_matches_the_latest_probability(self):
        org_id, client, case_id = self._setup("RecTier")
        self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        second = self._rank(client, case_id)
        _, recs = self._report(client, case_id)
        for name, r in second.items():
            p = r["predicted_probability"]
            tier = "high" if p >= 0.8 else "moderate" if p >= 0.4 else "low"
            self.assertIn(f"{tier} model-estimated probability", recs[name], name)

    def test_an_interrupted_ranking_that_stored_a_prediction_but_no_recommendation_is_still_consistent(self):
        """A prediction and its recommendation are two separate writes. If a ranking run were interrupted
        between them, the newest stored recommendation would belong to an EARLIER run. The report must
        still describe the latest prediction -- which is why it builds the text from that prediction
        instead of reading any stored recommendation."""
        org_id, client, case_id = self._setup("RecInterrupted")
        self._rank(client, case_id)
        self.assertEqual(self.upload(client, case_id, "second.csv", LOW).status_code, 201)
        second = self._rank(client, case_id)
        from backend.app.config.database import db_transaction
        from sqlalchemy import text
        with db_transaction() as conn:  # simulate the run dying after the predictions were written
            conn.execute(text(
                "DELETE FROM recommended_experiments WHERE id IN (SELECT MAX(r.id) FROM recommended_experiments r "
                "JOIN candidate_substitutes c ON c.id = r.candidate_id WHERE c.change_case_id = :cid GROUP BY r.candidate_id)"),
                {"cid": case_id})
        ranked, recs = self._report(client, case_id)
        for name in ranked:
            latest = self._pct(second[name]["predicted_probability"])
            self.assertEqual((ranked[name], self._estimate_in(recs[name])), (latest, latest), name)

    def test_a_candidate_with_a_prediction_never_gets_the_placeholder_text(self):
        org_id, client, case_id = self._setup("RecPlaceholder")
        self._rank(client, case_id)
        _, recs = self._report(client, case_id)
        for name, text in recs.items():
            self.assertNotIn("No experiment recommended yet", text, name)


if __name__ == "__main__":
    unittest.main()
