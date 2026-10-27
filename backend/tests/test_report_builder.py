"""
Real, executable test for the .docx report builder -- python-docx is
available in the sandbox and report_builder.py has no DB/web-framework
dependency (see its module docstring).
Run with: python3 -m unittest backend.tests.test_report_builder -v
"""
import os
import unittest

from docx import Document

from backend.app.services.report_builder import (
    build_project_report, build_qualification_report, format_predicted_value,
)


class TestReportBuilder(unittest.TestCase):
    def test_builds_a_real_readable_docx(self):
        project = {
            "name": "Adhesive Optimization", "target_metric": "bond_strength_MPa",
            "direction": "maximize", "target_value": 20.0,
        }
        backtest = {
            "actual_order_evals": 33, "engine_order_evals": 12,
            "reduction_pct": 63.6, "n_historical_rows": 60,
        }
        recommendations = [{
            "predicted_value": 21.4, "feasibility_probability": 0.91,
            "features": {"cure_temp_C": 137.0, "cure_time_min": 38.0},
        }]

        path = build_project_report(project, backtest, recommendations, 60, "Acme Coatings")
        self.assertTrue(os.path.exists(path))

        doc = Document(path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Adhesive Optimization", full_text)
        self.assertIn("Acme Coatings", full_text)
        os.remove(path)

    def test_handles_backtest_error_gracefully(self):
        project = {"name": "P", "target_metric": "y", "direction": "maximize", "target_value": None}
        backtest = {"error": "Not enough data."}
        path = build_project_report(project, backtest, [], 2, "Acme")
        self.assertTrue(os.path.exists(path))
        doc = Document(path)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Not enough data.", full_text)
        os.remove(path)


class TestQualificationReportModelEvidence(unittest.TestCase):
    """Priority 7B: model_quality/uncertainty_calibration must be surfaced
    honestly, never claimed unconditionally. The 'metrics unavailable'
    branch (n < 6 historical rows) can't be reached through the live API
    -- ranking itself requires >= 8 distinct rows (MIN_HISTORICAL_ROWS_
    FOR_PREDICTION), above compute_model_metrics'/compute_uncertainty_
    calibration's own n < 6 threshold -- so it's tested directly here,
    the same way test_handles_backtest_error_gracefully above already
    tests build_project_report's own error branch directly."""

    def _case(self):
        return {"name": "PFHxA replacement", "trigger_type": "regulatory_restriction",
                "restricted_substance": "PFHxA"}

    def _docx_text(self, path):
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        os.remove(path)
        return text

    def test_states_actual_metrics_when_available(self):
        ranked_results = [{
            "candidate_name": "Resin A", "predicted_probability": 0.8, "uncertainty_std": 0.05,
            "recommended_experiment": "Run a test.", "domain_coverage": None,
        }]
        model_quality = {
            "r2_score": 0.91, "mean_absolute_error": 1.2, "prediction_accuracy_pct": 86.7,
            "accuracy_tolerance_band": 2.5, "n_folds_used": 5, "n_points_scored": 15,
        }
        uncertainty_calibration = {
            "sigma_error_correlation": 0.42, "within_1sigma_pct": 71.4,
            "within_2sigma_pct": 92.9, "n_points_scored": 14,
        }
        path = build_qualification_report(
            self._case(), ranked_results, 15, "Acme Coatings",
            data_quality={"status": "valid", "dataset_id": 1, "row_count": 15,
                          "distinct_rows": 15, "duplicate_rows": 0, "constant_columns": []},
            model_quality=model_quality, uncertainty_calibration=uncertainty_calibration,
        )
        text = self._docx_text(path)
        self.assertIn("Model performance evidence", text)
        self.assertIn("R\u00b2 = 0.91", text)
        self.assertIn("mean absolute error = 1.2", text)
        self.assertIn("= 0.42", text)
        self.assertIn("71.4% of held-out points fell within 1", text)
        self.assertNotIn("with calibrated uncertainty -- not a heuristic score", text)

    def test_states_unavailable_when_metrics_have_errors(self):
        ranked_results = [{
            "candidate_name": "Resin A", "predicted_probability": 0.8, "uncertainty_std": 0.05,
            "recommended_experiment": "Run a test.", "domain_coverage": None,
        }]
        model_quality = {"error": "Need at least 6 historical rows to cross-validate model quality."}
        path = build_qualification_report(
            self._case(), ranked_results, 5, "Acme Coatings",
            data_quality={"status": "valid", "dataset_id": 1, "row_count": 5,
                          "distinct_rows": 5, "duplicate_rows": 0, "constant_columns": []},
            model_quality=model_quality, uncertainty_calibration=None,
        )
        text = self._docx_text(path)
        self.assertIn("Model performance evidence", text)
        self.assertIn("could not be computed", text)
        self.assertIn(model_quality["error"], text)
        # must never assert a calibration claim it has no evidence for
        self.assertNotIn("with calibrated uncertainty -- not a heuristic score", text)
        self.assertNotIn("R\u00b2 =", text)

    def test_no_model_evidence_section_when_metrics_not_provided(self):
        # Backward-compatible default: omitting model_quality entirely
        # (e.g. the insufficient-evidence report path) must not render an
        # empty or misleading evidence section.
        ranked_results = []
        path = build_qualification_report(self._case(), ranked_results, 6, "Acme Coatings",
                                           data_quality={"status": "insufficient", "dataset_id": 1,
                                                          "row_count": 6, "distinct_rows": 6,
                                                          "duplicate_rows": 0, "constant_columns": []})
        text = self._docx_text(path)
        self.assertNotIn("Model performance evidence", text)
        self.assertIn("Insufficient evidence to establish predictive performance.", text)


if __name__ == "__main__":
    unittest.main()

class TestQualificationReportDecisionSupport(unittest.TestCase):
    """C1: 'Model-estimated probability' wording and the 'Decision support'
    column/explanation. Table cells are read too (the 7B helper above only
    reads paragraphs)."""

    def _case(self):
        return {"name": "PFHxA replacement", "trigger_type": "regulatory_restriction",
                "restricted_substance": "PFHxA"}

    def _text(self, path):
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        os.remove(path)
        return "\n".join(parts)

    def _dq(self):
        return {"status": "valid", "dataset_id": 1, "row_count": 15,
                "distinct_rows": 15, "duplicate_rows": 0, "constant_columns": []}

    def _outside_result(self):
        from backend.app.services import change_case_rules as rules
        cov = rules.classify_candidate_domain_coverage({"viscosity": 500}, {"viscosity": (450, 464)})
        return {
            "candidate_name": "Far", "predicted_probability": 1.0, "uncertainty_std": 2.261,
            "recommended_experiment": "Run a test.", "domain_coverage": cov,
            "decision_support": rules.derive_decision_support(cov),
        }

    def test_outside_candidate_shows_model_estimate_and_requires_validation(self):
        path = build_qualification_report(self._case(), [self._outside_result()], 15, "Acme",
                                          data_quality=self._dq())
        text = self._text(path)
        self.assertIn("Model-estimated probability", text)
        self.assertIn("Decision support", text)
        self.assertIn("Requires validation", text)
        self.assertIn("Outside historical range", text)
        self.assertIn("100%", text)  # raw probability shown unchanged
        self.assertIn("2.261", text)
        self.assertIn("not a validated qualification probability", text)
        self.assertIn("'Within data range' does not mean qualified", text)

    def test_old_wording_removed(self):
        path = build_qualification_report(self._case(), [self._outside_result()], 15, "Acme",
                                          data_quality=self._dq())
        text = self._text(path)
        self.assertNotIn("Predicted qualification probability", text)
        self.assertNotIn("confidence-scored", text)
        self.assertIn("ranked shortlist ordered by model-estimated probability", text)

    def test_results_without_decision_support_omit_the_column(self):
        r = self._outside_result()
        r["decision_support"] = None
        path = build_qualification_report(self._case(), [r], 15, "Acme", data_quality=self._dq())
        text = self._text(path)
        self.assertNotIn("Decision support", text)
        self.assertIn("Model-estimated probability", text)


class TestPredictedValueInTheQualificationReport(unittest.TestCase):
    """Level 1: the ranked table carries the predicted target-property value beside the requirement,
    direction and unit; a value that was never recorded is labelled, not blanked or invented."""

    CASE = {"name": "PFHxA replacement", "trigger_type": "regulatory_restriction", "restricted_substance": "PFHxA"}
    CTX = {
        "change_case_id": 7, "generated_at_utc": "2026-10-13 10:00:00", "dataset": None, "stale_notice": None,
        "ingestion_review": None, "model_versions": ["gp-rank-v1"], "ranked_at_utc": None, "candidates": [],
        "outcomes": [],
        "spec": {"feature_columns": ["a"], "target_metric": "salt_spray_hours", "target_value": 500.0,
                 "direction": "maximize", "units": {"salt_spray_hours": "h"}},
    }

    def _text(self, path):
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        os.remove(path)
        return "\n".join(parts)

    def _result(self, cid, name, value, prob=0.7, rank=1):
        return {"candidate_id": cid, "rank": rank, "candidate_name": name, "predicted_value": value,
                "predicted_probability": prob, "uncertainty_std": 12.345,
                "recommended_experiment": f"Run a physical qualification test for '{name}'."}

    def test_value_threshold_direction_unit_and_uncertainty_are_shown(self):
        path = build_qualification_report(self.CASE, [self._result(1, "Resin A", 523.678)], 12, "Acme",
                                          audit_context=self.CTX)
        text = self._text(path)
        self.assertIn("Requirement: salt_spray_hours (h) at or above 500.0 (direction: maximize)", text)
        self.assertIn("Predicted salt_spray_hours (h) (model estimate)", text)
        self.assertIn("523.7", text)
        self.assertIn("Uncertainty (std, h)", text)
        self.assertIn("12.345", text)
        self.assertIn("MODEL ESTIMATES, not measured physical results", text)

    def test_minimize_direction_is_stated(self):
        ctx = {**self.CTX, "spec": {**self.CTX["spec"], "direction": "minimize"}}
        text = self._text(build_qualification_report(self.CASE, [self._result(1, "A", 10.0)], 12, "Acme",
                                                     audit_context=ctx))
        self.assertIn("at or below 500.0 (direction: minimize)", text)

    def test_missing_value_is_labelled_not_blank_and_not_invented(self):
        text = self._text(build_qualification_report(
            self.CASE, [self._result(1, "Old", None, rank=2), self._result(2, "New", 410.0, prob=0.9, rank=1)],
            12, "Acme", audit_context=self.CTX))
        self.assertIn("Not recorded", text)
        self.assertIn("No value has been estimated or back-filled", text)
        self.assertIn("410", text)

    def test_note_about_missing_values_only_appears_when_needed(self):
        text = self._text(build_qualification_report(self.CASE, [self._result(1, "A", 410.0)], 12, "Acme",
                                                     audit_context=self.CTX))
        self.assertNotIn("Not recorded", text)
        self.assertNotIn("back-filled", text)

    def test_rank_order_note_is_present_with_ranks(self):
        text = self._text(build_qualification_report(self.CASE, [self._result(1, "A", 410.0)], 12, "Acme",
                                                     audit_context=self.CTX))
        self.assertIn("then by candidate ID", text)

    def test_results_without_the_new_fields_still_render_as_before(self):
        legacy = [{"candidate_name": "Old", "predicted_probability": 0.5, "uncertainty_std": 1.0,
                   "recommended_experiment": "Run it."}]
        doc = Document(build_qualification_report(self.CASE, legacy, 12, "Acme"))
        table = doc.tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells],
                         ["Candidate", "Model-estimated probability", "Uncertainty (std)"])

    def test_format_predicted_value(self):
        self.assertEqual(format_predicted_value(523.678), "523.7")
        self.assertEqual(format_predicted_value(12345.6), "12350")
        self.assertEqual(format_predicted_value(100000.0), "100000")
        self.assertEqual(format_predicted_value(0.0000123456), "0.00001235")
        self.assertEqual(format_predicted_value(-3.14159), "-3.142")
        self.assertEqual(format_predicted_value(0.0), "0")
        for v in (523.678, 12345.6, 1e9, 1e-9, -77.7):
            self.assertNotIn("e", format_predicted_value(v).lower())
