"""
Evidence extraction: turns an uploaded customer file into candidate TABLES for
human review. It extracts; it never decides.

Supported: .csv, .xlsx, .docx (Word tables), .pdf (text-based PDFs with
tables). NOT supported, and refused with an explanation rather than guessed
at: .xls / .xlsm / .doc (re-save as .xlsx / .docx), password-protected files,
scanned/image PDFs (no OCR), tables nested inside tables, charts and images.

Every table carries where it came from (file kind, sheet / table / page) and
every row carries its source row number, so a value in the validated dataset
can be traced back to a spot in the customer's file.

What this module will NOT do silently:
- pick one of several tables (the reviewer selects; see evidence_rules flags)
- decide which row is the header (default: the first non-empty row, shown to
  the reviewer, who can name another via `header_row`)
- convert or interpret units / conditions (that is evidence_rules' job)
"""
import csv
import io
import re

from backend.app.config.settings import settings
from backend.app.schemas.errors import ValidationError
from backend.app.services import evidence_rules

SUPPORTED_EXTENSIONS = {".csv": "csv", ".xlsx": "xlsx", ".docx": "docx", ".pdf": "pdf"}
ACCEPTED_TEXT = ".csv, .xlsx, .docx, .pdf"

MAX_TABLES = 40
MAX_XLSX_SHEETS = 20
MAX_PDF_PAGES = 100
MAX_COLUMNS = 60
PREVIEW_ROWS = 40

_QUOTED_TEXT_IN_FORMAT = re.compile(r'"[^"]+"')


def file_kind(filename: str) -> str:
    name = (filename or "").lower()
    for ext, kind in SUPPORTED_EXTENSIONS.items():
        if name.endswith(ext):
            return kind
    hints = {
        ".xls": "Legacy .xls files are not supported -- re-save the workbook as .xlsx.",
        ".xlsm": "Macro-enabled workbooks are not supported -- save a copy as .xlsx (values only).",
        ".doc": "Legacy .doc files are not supported -- re-save the document as .docx.",
    }
    for ext, hint in hints.items():
        if name.endswith(ext):
            raise ValidationError(f"{hint} Accepted file types: {ACCEPTED_TEXT}.")
    raise ValidationError(f"Unsupported file type. Accepted file types: {ACCEPTED_TEXT}.")


def _table(ref, kind, location, rows, caption="", confidence="native", column_warnings=None, extra=None):
    t = {
        "ref": ref,
        "kind": kind,
        "location": location,
        "caption": (caption or "").strip(),
        "confidence": confidence,
        "rows": rows,  # [{"src_row": int, "cells": [str, ...]}]
        "column_warnings": column_warnings or [],
        "warnings": [],
    }
    if extra:
        t.update(extra)
    return t


def _clean(cell) -> str:
    if cell is None:
        return ""
    return re.sub(r"\s+", " ", str(cell)).strip()


def _trim_columns(rows):
    """Drop trailing columns that are empty in every row, pad the rest to a
    rectangle, and cap the width."""
    width = 0
    for r in rows:
        for i in range(len(r["cells"]) - 1, -1, -1):
            if r["cells"][i] != "":
                width = max(width, i + 1)
                break
    width = min(width, MAX_COLUMNS)
    for r in rows:
        cells = r["cells"][:width]
        cells += [""] * (width - len(cells))
        r["cells"] = cells
    return width


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def _csv_tables(file_bytes: bytes):
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValidationError("Could not read file as UTF-8 text. Please export as a standard CSV.")
    reader = csv.reader(io.StringIO(text))
    rows = []
    for raw in reader:
        if not raw:  # a truly blank line (",,," rows are kept and counted as blank data rows later)
            continue
        if len(rows) > settings.MAX_UPLOAD_ROWS + 1:
            raise ValidationError(f"Too many rows (max {settings.MAX_UPLOAD_ROWS}).")
        rows.append({"src_row": reader.line_num, "cells": [str(c) for c in raw]})
    if not rows:
        raise ValidationError("Could not find a header row in the CSV.")
    return [_table("csv", "csv", "CSV file", rows)], []


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
def _xlsx_cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return repr(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return _clean(v)


def _xlsx_tables(file_bytes: bytes):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        raise ValidationError(
            "Could not read this .xlsx file. It may be corrupted, password-protected, or not a real "
            "Excel workbook. Password-protected files are not supported -- remove the password and re-upload."
        )
    tables, warnings = [], []
    try:
        sheets = list(wb.worksheets)
        if len(sheets) > MAX_XLSX_SHEETS:
            raise ValidationError(f"Workbook has {len(sheets)} sheets (limit {MAX_XLSX_SHEETS}). Split it or remove unused sheets.")
        for ws in sheets:
            rows, fmt_by_col = [], {}
            truncated = False
            for r_idx, row in enumerate(ws.iter_rows(), start=1):
                if len(rows) > settings.MAX_UPLOAD_ROWS + 1:
                    truncated = True
                    break
                cells = []
                for c_idx, cell in enumerate(row):
                    v = getattr(cell, "value", None)
                    cells.append(_xlsx_cell_text(v))
                    fmt = getattr(cell, "number_format", None)
                    if isinstance(v, (int, float)) and not isinstance(v, bool) and fmt and _QUOTED_TEXT_IN_FORMAT.search(fmt):
                        fmt_by_col.setdefault(c_idx, set()).add(fmt)
                if all(c == "" for c in cells):
                    continue
                rows.append({"src_row": r_idx, "cells": cells})
            if truncated:
                raise ValidationError(f"Sheet '{ws.title}' has too many rows (max {settings.MAX_UPLOAD_ROWS}).")
            if not rows:
                continue
            width = _trim_columns(rows)
            hidden = getattr(ws, "sheet_state", "visible") != "visible"
            colw = [
                {
                    "col_index": ci,
                    "code": "XLSX_NUMBER_FORMAT_HIDES_TEXT",
                    "message": (
                        f"Cells in this column display extra text through their number format "
                        f"({', '.join(sorted(fmts))}) that is NOT part of the stored value (for example a unit). "
                        "The value will be read as a bare number; declare its unit explicitly."
                    ),
                }
                for ci, fmts in sorted(fmt_by_col.items()) if ci < width
            ]
            t = _table(f"xlsx:{ws.title}", "xlsx", f"Sheet '{ws.title}'" + (" (hidden sheet)" if hidden else ""),
                       rows, column_warnings=colw)
            if hidden:
                t["warnings"].append("This sheet is hidden in the workbook; confirm it is current before relying on it.")
            tables.append(t)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    if tables:
        warnings.append(
            "Excel values are read as stored (cached results of formulas, merged cells empty except the "
            "top-left cell). Charts, comments and cell colours are ignored."
        )
    return tables, warnings


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _docx_tables(file_bytes: bytes):
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(io.BytesIO(file_bytes))
    except Exception:
        raise ValidationError("Could not read this .docx file. It may be corrupted or password-protected.")
    tables, warnings = [], []
    last_text = ""
    t_idx = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            try:
                txt = Paragraph(child, doc).text.strip()
            except Exception:
                txt = ""
            if txt:
                last_text = txt
        elif tag == "tbl":
            t_idx += 1
            if len(tables) >= MAX_TABLES:
                raise ValidationError(f"Document has more than {MAX_TABLES} tables. Split it or extract the relevant tables.")
            try:
                tbl = Table(child, doc)
                rows, nested = [], False
                for r_i, row in enumerate(tbl.rows, start=1):
                    seen, cells = set(), []
                    for cell in row.cells:
                        key = id(cell._tc)
                        if key in seen:  # horizontally merged cell: keep the text once
                            cells.append("")
                            continue
                        seen.add(key)
                        if cell.tables:
                            nested = True
                        cells.append(_clean(cell.text))
                    if all(c == "" for c in cells):
                        continue
                    rows.append({"src_row": r_i, "cells": cells})
            except Exception:
                warnings.append(f"Table {t_idx} could not be read and was skipped.")
                continue
            if not rows:
                continue
            _trim_columns(rows)
            t = _table(f"docx:t{t_idx}", "docx", f"Table {t_idx}", rows, caption=last_text[:200])
            if nested:
                t["warnings"].append("This table contains nested tables, which are ignored.")
            tables.append(t)
    if tables:
        warnings.append("Word tables are read as text. Text, headings and notes outside tables are NOT parsed "
                        "(the nearest preceding paragraph is shown as context only).")
    return tables, warnings


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _pdf_grid(raw_table):
    rows = []
    for i, r in enumerate(raw_table, start=1):
        cells = [_clean(c) for c in (r or [])]
        if all(c == "" for c in cells):
            continue
        rows.append({"src_row": i, "cells": cells})
    return rows


def _pdf_tables(file_bytes: bytes):
    try:
        import pdfplumber
    except Exception:  # pragma: no cover
        raise ValidationError("PDF support is not installed on this server.")
    tables, warnings = [], []
    try:
        pdf = pdfplumber.open(io.BytesIO(file_bytes))
    except Exception:
        raise ValidationError(
            "Could not read this PDF. It may be corrupted or password-protected. "
            "Password-protected PDFs are not supported -- remove the password and re-upload."
        )
    try:
        n_pages = len(pdf.pages)
        if n_pages > MAX_PDF_PAGES:
            raise ValidationError(
                f"This PDF has {n_pages} pages (limit {MAX_PDF_PAGES}). Extract the pages that contain the "
                "relevant tables and upload those."
            )
        total_text = 0
        page_context = {}
        for pno, page in enumerate(pdf.pages, start=1):
            try:
                txt = page.extract_text() or ""
            except Exception:
                txt = ""
            total_text += len(txt.strip())
            page_context[pno] = re.sub(r"\s+", " ", txt).strip()[:200]
            try:
                found = page.extract_tables()
            except Exception:
                found = []
            for tno, raw in enumerate(found, start=1):
                rows = _pdf_grid(raw)
                if len(rows) < 2:
                    continue
                if len(tables) >= MAX_TABLES:
                    raise ValidationError(f"PDF has more than {MAX_TABLES} tables. Extract the relevant pages and upload those.")
                _trim_columns(rows)
                tables.append(_table(f"pdf:p{pno}t{tno}", "pdf", f"Page {pno}, table {tno}", rows,
                                     caption=page_context[pno], confidence="ruled"))
        if not tables and total_text >= 20:
            # No ruled tables anywhere: try layout-inferred tables, clearly marked lower confidence.
            for pno, page in enumerate(pdf.pages, start=1):
                try:
                    found = page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
                except Exception:
                    found = []
                for tno, raw in enumerate(found, start=1):
                    rows = _pdf_grid(raw)
                    if len(rows) < 3:
                        continue
                    width = _trim_columns(rows)
                    if width < 2:
                        continue
                    if len(tables) >= MAX_TABLES:
                        break
                    t = _table(f"pdf:p{pno}t{tno}", "pdf", f"Page {pno}, table {tno} (inferred from text layout)",
                               rows, caption=page_context.get(pno, ""), confidence="text_layout")
                    t["warnings"].append(
                        "No ruled table was found; this table was INFERRED from text alignment and may have "
                        "merged or split cells. Check every value against the source page."
                    )
                    tables.append(t)
        if total_text < 20:
            warnings.append(
                "This PDF has no extractable text layer (it looks scanned or image-only). OCR is not supported: "
                "transcribe the relevant values into a spreadsheet and upload that."
            )
        elif tables:
            warnings.append(
                "PDF tables are extracted from the text layer. A table that continues across pages appears as "
                "separate tables; values in figures or images are not extracted."
            )
    finally:
        try:
            pdf.close()
        except Exception:
            pass
    return tables, warnings


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------
def list_tables(filename: str, file_bytes: bytes) -> dict:
    """-> {"kind", "tables": [table...], "warnings": [str...]}. Raises
    ValidationError (user-safe message) for unsupported/unreadable files."""
    kind = file_kind(filename)
    if kind == "csv":
        tables, warnings = _csv_tables(file_bytes)
    elif kind == "xlsx":
        tables, warnings = _xlsx_tables(file_bytes)
    elif kind == "docx":
        tables, warnings = _docx_tables(file_bytes)
    else:
        tables, warnings = _pdf_tables(file_bytes)
    return {"kind": kind, "tables": tables, "warnings": warnings}


def default_header_row(table: dict) -> int:
    """The source row number of the table's first non-empty row."""
    return table["rows"][0]["src_row"]


def table_summary(table: dict, spec_required=None) -> dict:
    """JSON-safe summary for the review screen: location, a preview, and how
    well its first row matches the spec's required columns (display only --
    nothing is applied from this)."""
    rows = table["rows"]
    header = rows[0]["cells"] if rows else []
    out = {
        "ref": table["ref"],
        "kind": table["kind"],
        "location": table["location"],
        "caption": table["caption"],
        "confidence": table["confidence"],
        "n_rows": len(rows),
        "n_cols": len(header),
        "default_header_row": default_header_row(table) if rows else None,
        "preview": rows[:PREVIEW_ROWS],
        "preview_truncated": len(rows) > PREVIEW_ROWS,
        "warnings": list(table["warnings"]),
        "column_warnings": table["column_warnings"],
    }
    if spec_required:
        headers = [h.strip() for h in header]
        matches = {}
        for canon in spec_required:
            cands = evidence_rules.propose_header_matches(canon, headers)
            matches[canon] = {"best": cands[0]["header"] if cands else None,
                              "confidence": cands[0]["confidence"] if cands else None,
                              "n_candidates": len(cands)}
        out["matches"] = matches
        out["matched_required"] = sum(1 for m in matches.values() if m["best"] is not None and m["n_candidates"] == 1)
        out["total_required"] = len(spec_required)
    return out


def _normalize_refs(table_ref):
    if table_ref is None or table_ref == "" or table_ref == []:
        return []
    if isinstance(table_ref, str):
        return [table_ref]
    return list(table_ref)


def load_selection(filename: str, file_bytes: bytes, options: dict) -> dict:
    """Resolves WHICH table(s) feed the dataset, strictly from reviewer input.

    Returns {"headers", "data_rows", "row_refs", "source", "flags", "listing"}.
    When selection cannot be resolved without guessing, `headers` is None and
    `flags` holds blocking review items (never a silent default). Unreadable /
    unsupported files raise ValidationError."""
    options = options or {}
    listing = list_tables(filename, file_bytes)
    kind, tables = listing["kind"], listing["tables"]
    flags = []
    result = {"headers": None, "data_rows": [], "row_refs": [], "source": None, "flags": flags, "listing": listing}

    if not tables:
        flags.append(evidence_rules.make_flag(
            evidence_rules.BLOCKING, "NO_TABLES_FOUND",
            "No tables were found in this file. " + (" ".join(listing["warnings"]) or
            "Transcribe the relevant values into a spreadsheet and upload that."),
        ))
        return result

    refs = _normalize_refs(options.get("table_ref"))
    by_ref = {t["ref"]: t for t in tables}
    if not refs:
        if len(tables) == 1:
            chosen = [tables[0]]
            flags.append(evidence_rules.make_flag(
                evidence_rules.WARNING, "TABLE_AUTO_SELECTED",
                f"The file contains exactly one table ({tables[0]['location']}); it was used. Confirm it is the right one.",
            ))
        else:
            desc = "; ".join(f"{t['ref']} = {t['location']} ({len(t['rows'])} rows)" for t in tables[:12])
            flags.append(evidence_rules.make_flag(
                evidence_rules.BLOCKING, "TABLE_SELECTION_REQUIRED",
                f"The file contains {len(tables)} tables and the software will not choose between them. "
                f"Select the table(s) to use via table_ref. Found: {desc}.",
            ))
            return result
    else:
        missing = [r for r in refs if r not in by_ref]
        if missing:
            flags.append(evidence_rules.make_flag(
                evidence_rules.BLOCKING, "INVALID_OPTIONS",
                f"table_ref {', '.join(repr(m) for m in missing)} is not a table in this file.",
            ))
            return result
        chosen = [by_ref[r] for r in refs]

    header_row = options.get("header_row")
    if header_row is not None and len(chosen) > 1:
        flags.append(evidence_rules.make_flag(
            evidence_rules.BLOCKING, "INVALID_OPTIONS",
            "header_row can only be given when exactly one table is selected; with several tables each "
            "table's first row must be its header row.",
        ))
        return result

    per_table = []
    for t in chosen:
        hr = header_row if header_row is not None else default_header_row(t)
        idx = next((i for i, r in enumerate(t["rows"]) if r["src_row"] == hr), None)
        if idx is None:
            flags.append(evidence_rules.make_flag(
                evidence_rules.BLOCKING, "INVALID_OPTIONS",
                f"header_row {hr} is not a non-empty row of {t['location']}.",
            ))
            return result
        headers = [c.strip() for c in t["rows"][idx]["cells"]]
        per_table.append((t, hr, headers, t["rows"][idx + 1:]))

    first_headers = per_table[0][2]
    for t, hr, headers, _ in per_table[1:]:
        if headers != first_headers:
            flags.append(evidence_rules.make_flag(
                evidence_rules.BLOCKING, "HEADERS_DIFFER_ACROSS_TABLES",
                f"{t['location']} has different column headers from {per_table[0][0]['location']}; tables can only be "
                "combined when their header rows are identical. Select one table, or combine them in a spreadsheet first.",
            ))
            return result

    data_rows, row_refs = [], []
    for t, hr, headers, rows in per_table:
        for r in rows:
            if len(data_rows) >= settings.MAX_UPLOAD_ROWS:
                raise ValidationError(f"Too many rows (max {settings.MAX_UPLOAD_ROWS}).")
            cells = r["cells"]
            data_rows.append({h: (cells[i] if i < len(cells) else None) for i, h in enumerate(headers)})
            row_refs.append({"file": filename, "table": t["ref"], "location": t["location"], "row": r["src_row"]})

    for t, hr, headers, _ in per_table:
        for w in t["warnings"]:
            flags.append(evidence_rules.make_flag(evidence_rules.WARNING, "TABLE_WARNING", f"{t['location']}: {w}"))
        for cw in t["column_warnings"]:
            ci = cw["col_index"]
            col = headers[ci] if ci < len(headers) else None
            if col:
                flags.append(evidence_rules.make_flag(evidence_rules.WARNING, cw["code"], f"{t['location']}: {cw['message']}", column=col))
    for w in listing["warnings"]:
        flags.append(evidence_rules.make_flag(evidence_rules.WARNING, "EXTRACTION_NOTE", w))

    result.update({
        "headers": first_headers,
        "data_rows": data_rows,
        "row_refs": row_refs,
        "source": {
            "file": filename,
            "kind": kind,
            "tables": [
                {"ref": t["ref"], "location": t["location"], "caption": t["caption"], "confidence": t["confidence"],
                 "header_row": hr, "data_rows": len(rows)}
                for t, hr, _, rows in per_table
            ],
        },
    })
    return result
